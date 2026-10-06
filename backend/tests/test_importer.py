"""Unit tests for spreadsheet importer service (parsing, conversion, mapping, template)."""

import csv
import datetime
import io
from unittest.mock import MagicMock
import zipfile
import openpyxl
import pytest

from app.errors import PollValidationError
from app.models import GroupField
from app.services.importer import (
    convert_cell_value,
    generate_suggested_mapping,
    generate_template_file,
    parse_raw_spreadsheet,
    validate_file_type_and_size,
)


def make_xlsx(rows: list[list]) -> bytes:
    """Helper to create an in-memory xlsx file bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    for r in rows:
        ws.append(r)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def make_csv(rows: list[list], delimiter: str = ",", encoding: str = "utf-8") -> bytes:
    """Helper to create in-memory csv bytes."""
    out = io.StringIO()
    writer = csv.writer(out, delimiter=delimiter)
    for r in rows:
        writer.writerow(r)
    return out.getvalue().encode(encoding)


# ---------------------------------------------------------------------------
# 1. File Type and Size Validation
# ---------------------------------------------------------------------------

def test_file_validation_size_limit():
    """File size exceeding 5 MB is rejected."""
    huge_content = b"x" * (5 * 1024 * 1024 + 1)
    with pytest.raises(PollValidationError) as exc:
        validate_file_type_and_size("roster.csv", huge_content)
    assert "5 MB" in exc.value.message


def test_file_validation_legacy_excel_rejection():
    """.xls and .xlsm files are rejected with specific advice."""
    for bad_name in ["roster.xls", "ROSTER.XLS", "data.xlsm", "DATA.XLSM"]:
        with pytest.raises(PollValidationError) as exc:
            validate_file_type_and_size(bad_name, b"PKdummy")
        assert "Please save the file as .xlsx or .csv." in exc.value.message


def test_file_validation_unsupported_extensions():
    """Other non-spreadsheet extensions are rejected with clear message."""
    for bad_name in ["roster.txt", "sheet.pdf", "data.json", "file"]:
        with pytest.raises(PollValidationError) as exc:
            validate_file_type_and_size(bad_name, b"PKdummy")
        assert "Unsupported file type" in exc.value.message


def test_file_validation_fake_xlsx():
    """A .xlsx file not starting with PK zip signature is rejected."""
    fake_content = b"Not a zip file"
    with pytest.raises(PollValidationError) as exc:
        validate_file_type_and_size("roster.xlsx", fake_content)
    assert "Invalid file content" in exc.value.message


def test_file_validation_uncompressed_size_limit():
    """Zip with uncompressed size > 50 MB is rejected before loading."""
    mock_info = MagicMock()
    mock_info.file_size = 50 * 1024 * 1024 + 10
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(zipfile.ZipFile, "infolist", lambda self: [mock_info])
        # Valid zip content that starts with PK
        bio = io.BytesIO()
        with zipfile.ZipFile(bio, "w") as zf:
            zf.writestr("test.xml", b"x")
        content = bio.getvalue()

        with pytest.raises(PollValidationError) as exc:
            validate_file_type_and_size("roster.xlsx", content)
        assert "50 MB" in exc.value.message


# ---------------------------------------------------------------------------
# 2. Cell Conversion
# ---------------------------------------------------------------------------

def test_convert_cell_value_rules():
    """Verify conversion rules for None, strings, numbers, dates, booleans."""
    # None -> blank string
    assert convert_cell_value(None) == ""

    # Strings -> trimmed
    assert convert_cell_value("  hello  ") == "hello"

    # Booleans -> TRUE / FALSE
    assert convert_cell_value(True) == "TRUE"
    assert convert_cell_value(False) == "FALSE"

    # Dates and datetimes -> YYYY-MM-DD
    d = datetime.date(2025, 3, 14)
    dt = datetime.datetime(2025, 3, 14, 15, 30, 0)
    assert convert_cell_value(d) == "2025-03-14"
    assert convert_cell_value(dt) == "2025-03-14"

    # Whole-number float 21.0 -> "21" for text field, 21 for number field
    assert convert_cell_value(21.0, target_field_type="text") == "21"
    assert convert_cell_value(21.0, target_field_type=None) == "21"
    assert convert_cell_value(21.0, target_field_type="number") == 21

    # Integer 21001 -> "21001" for text field, 21001 for number field
    assert convert_cell_value(21001, target_field_type="text") == "21001"
    assert convert_cell_value(21001, target_field_type="number") == 21001

    # Non-integer float 3.14 -> "3.14" for text, 3.14 for number
    assert convert_cell_value(3.14, target_field_type="text") == "3.14"
    assert convert_cell_value(3.14, target_field_type="number") == 3.14


# ---------------------------------------------------------------------------
# 3. Parsing Raw Spreadsheets (.xlsx & .csv)
# ---------------------------------------------------------------------------

def test_parse_raw_spreadsheet_xlsx_and_csv_parity():
    """xlsx and csv produce identical headers and data rows."""
    data = [
        ["Name", "Register No", "Department"],
        ["Alice", "21001", "CSE"],
        ["Bob", "21002", "ECE"],
    ]
    xlsx_bytes = make_xlsx(data)
    csv_bytes = make_csv(data)

    sheet_xlsx, headers_xlsx, rows_xlsx = parse_raw_spreadsheet("roster.xlsx", xlsx_bytes)
    sheet_csv, headers_csv, rows_csv = parse_raw_spreadsheet("roster.csv", csv_bytes)

    assert headers_xlsx == headers_csv == ["Name", "Register No", "Department"]
    assert len(rows_xlsx) == len(rows_csv) == 2
    assert rows_xlsx[0][1] == rows_csv[0][1] == ["Alice", "21001", "CSE"]
    assert rows_xlsx[1][1] == rows_csv[1][1] == ["Bob", "21002", "ECE"]


def test_parse_csv_with_bom_and_semicolon():
    """CSV with utf-8-sig BOM and semicolon delimiter is handled seamlessly."""
    raw = "\ufeffName;Dept\nAlice;CSE\nBob;ECE".encode("utf-8")
    _, headers, rows = parse_raw_spreadsheet("roster.csv", raw)
    assert headers == ["Name", "Dept"]
    assert len(rows) == 2
    assert rows[0][1] == ["Alice", "CSE"]


def test_parse_csv_cp1252_fallback():
    """CSV encoded with cp1252 falls back properly."""
    text = "Name,City\nRenée,Montréal\n"
    raw = text.encode("cp1252")
    _, headers, rows = parse_raw_spreadsheet("roster.csv", raw)
    assert headers == ["Name", "City"]
    assert rows[0][1] == ["Renée", "Montréal"]


def test_parse_spreadsheet_skips_blank_rows_and_preserves_row_numbers():
    """Leading blank rows and inner blank rows are skipped, preserving sheet row numbers."""
    rows = [
        [None, None],        # Row 1 (blank)
        ["", "  "],          # Row 2 (blank)
        ["Name", "Reg"],     # Row 3 (Header row!)
        [None, None],        # Row 4 (blank)
        ["Alice", "101"],    # Row 5 (Data row 1)
        ["", ""],            # Row 6 (blank)
        ["Bob", "102"],      # Row 7 (Data row 2)
    ]
    xlsx_bytes = make_xlsx(rows)
    _, headers, data_rows = parse_raw_spreadsheet("roster.xlsx", xlsx_bytes)
    assert headers == ["Name", "Reg"]
    assert len(data_rows) == 2
    assert data_rows[0][0] == 5
    assert data_rows[0][1] == ["Alice", "101"]
    assert data_rows[1][0] == 7
    assert data_rows[1][1] == ["Bob", "102"]


def test_parse_spreadsheet_blank_header_error():
    """A blank header column raises validation error naming the column."""
    rows = [
        ["Name", "", "Dept"],
        ["Alice", "101", "CSE"],
    ]
    with pytest.raises(PollValidationError) as exc:
        parse_raw_spreadsheet("roster.xlsx", make_xlsx(rows))
    assert "Blank header at column 2" in exc.value.message


def test_parse_spreadsheet_duplicate_header_error():
    """Duplicate header names (case-insensitive) raise validation error."""
    rows = [
        ["Name", "Dept", "dept"],
        ["Alice", "CSE", "CSE"],
    ]
    with pytest.raises(PollValidationError) as exc:
        parse_raw_spreadsheet("roster.xlsx", make_xlsx(rows))
    assert "Duplicate header 'dept' at column 3" in exc.value.message


def test_parse_spreadsheet_column_limit():
    """More than 60 columns is rejected naming the 60 limit."""
    headers = [f"Col{i}" for i in range(61)]
    rows = [headers, ["val"] * 61]
    with pytest.raises(PollValidationError) as exc:
        parse_raw_spreadsheet("roster.xlsx", make_xlsx(rows))
    assert "60 columns" in exc.value.message


def test_parse_spreadsheet_row_limit():
    """More than 2000 data rows is rejected naming the 2000 limit."""
    rows = [["Name"]] + [[f"Member {i}"] for i in range(2001)]
    with pytest.raises(PollValidationError) as exc:
        parse_raw_spreadsheet("roster.csv", make_csv(rows))
    assert "2000 data rows" in exc.value.message


# ---------------------------------------------------------------------------
# 4. Suggested Mapping
# ---------------------------------------------------------------------------

def test_suggested_mapping_rules():
    """Matches Name, Full name, Student name to 'name' and fields by name/key."""
    f_dept = MagicMock(spec=GroupField)
    f_dept.name = "Department"
    f_dept.key = "dept_key"

    f_reg = MagicMock(spec=GroupField)
    f_reg.name = "Register No"
    f_reg.key = "reg_no"

    fields = [f_dept, f_reg]

    headers = ["Full Name", "Register no", "department", "Unknown"]
    mapping = generate_suggested_mapping(headers, fields)

    assert mapping["0"] == "name"
    assert mapping["1"] == "reg_no"
    assert mapping["2"] == "dept_key"
    assert mapping["3"] == "skip"


# ---------------------------------------------------------------------------
# 5. Template Generation
# ---------------------------------------------------------------------------

def test_template_generation_xlsx_and_csv():
    """Template includes Name first, then fields in position order, with choice dropdown."""
    f1 = MagicMock(spec=GroupField)
    f1.name = "Department"
    f1.position = 2
    f1.field_type = "choice"
    f1.choices = ["CSE", "ECE", "MECH"]

    f2 = MagicMock(spec=GroupField)
    f2.name = "Register No"
    f2.position = 1
    f2.field_type = "text"
    f2.choices = None

    fields = [f1, f2]

    # 1. CSV Template
    csv_bytes, media_csv, fname_csv = generate_template_file(fields, "csv")
    assert media_csv.startswith("text/csv")
    assert fname_csv == "members-template.csv"
    assert csv_bytes.decode("utf-8").strip() == "Name,Register No,Department"

    # 2. XLSX Template
    xlsx_bytes, media_xlsx, fname_xlsx = generate_template_file(fields, "xlsx")
    assert "spreadsheetml.sheet" in media_xlsx
    assert fname_xlsx == "members-template.xlsx"

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert [cell.value for cell in ws[1]] == ["Name", "Register No", "Department"]
    # Check data validation exists
    assert len(ws.data_validations.dataValidation) == 1
    dv = ws.data_validations.dataValidation[0]
    assert dv.type == "list"
    assert '"CSE,ECE,MECH"' in dv.formula1
    assert "C2:C2001" in str(dv.sqref)
