"""Tests for Part FB: Deleting members, polls, and groups, bulk selection, roster export."""

from datetime import datetime, timezone
import io
import threading
import time
import uuid
import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker

from app.errors import PollNotFoundError, PollValidationError
from app.models import (
    Group,
    GroupField,
    Member,
    Poll,
    PollAnswer,
    PollField,
    PollIncludedField,
    PollOption,
    Vote,
    VoteHistory,
)
from app.services.fields import create_group_field
from app.services.identity import (
    add_members_bulk,
    claim_member,
    create_group_with_admin,
)
from app.services.polls import (
    PollOptionInput,
    cast_vote,
    create_poll,
    delete_poll,
    save_poll_answers,
)


def _setup_group_with_data(db_session, name="Group Alpha"):
    """Helper to create a group with fields, members, claimed tokens, a poll, votes, and answers."""
    group, admin_token = create_group_with_admin(db_session, name)

    # 2 group fields
    f1 = create_group_field(db_session, group.id, "Department", "text")
    f2 = create_group_field(db_session, group.id, "RegisterNo", "text")

    # 3 members
    members = add_members_bulk(
        db_session,
        group.id,
        members_data=[
            {"display_name": "Alice", "values": {f1.key: "CS", f2.key: "R101"}},
            {"display_name": "Bob", "values": {f1.key: "EE", f2.key: "R102"}},
            {"display_name": "=Charlie", "values": {f1.key: "+Lead", f2.key: "-R103"}},
        ],
    )

    # Claim Alice
    _, alice_token = claim_member(db_session, group.join_code, members[0].id)

    # Create a poll with poll fields and included fields
    poll = create_poll(
        db_session,
        group_id=group.id,
        name=f"Poll for {name}",
        allow_multiple=False,
        options=[
            PollOptionInput(label="Finished", role="target"),
            PollOptionInput(label="Working", role="in_progress"),
        ],
        included_field_ids=[f1.id, f2.id],
        poll_fields=[{"name": "Comments", "field_type": "text"}],
    )

    # Alice votes and answers
    opt_target = [o for o in poll.options if o.role == "target"][0]
    cast_vote(db_session, poll.id, members[0].id, opt_target.id)
    save_poll_answers(
        db_session,
        poll.id,
        members[0].id,
        values={poll.poll_fields[0].key: "All good"},
    )

    db_session.commit()

    return {
        "group": group,
        "admin_token": admin_token,
        "fields": [f1, f2],
        "members": members,
        "alice_token": alice_token,
        "poll": poll,
    }


# ---------------------------------------------------------------------------
# 1. Single Member Deletion (DELETE /api/v1/groups/{group_id}/members/{member_id})
# ---------------------------------------------------------------------------

def test_delete_member_cascades_and_invalidates_token(client, db_session):
    """Single member delete removes their votes, history, and answers, and leaves others untouched."""
    data = _setup_group_with_data(db_session, "Member Delete Group")
    group_id = data["group"].id
    admin_tok = data["admin_token"]
    alice_id = data["members"][0].id
    bob_id = data["members"][1].id
    alice_tok = data["alice_token"]

    # Verify Alice has votes and answers
    assert db_session.scalar(select(func.count(Vote.id)).where(Vote.member_id == alice_id)) == 1
    assert db_session.scalar(select(func.count(VoteHistory.id)).where(VoteHistory.member_id == alice_id)) == 1
    assert db_session.scalar(select(func.count(PollAnswer.id)).where(PollAnswer.member_id == alice_id)) == 1

    # Delete Alice
    res = client.delete(
        f"/api/v1/groups/{group_id}/members/{alice_id}",
        headers={"X-Admin-Token": admin_tok},
    )
    assert res.status_code == 204

    # Alice is gone
    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == alice_id)) == 0
    # Alice's votes, history, answers are gone
    assert db_session.scalar(select(func.count(Vote.id)).where(Vote.member_id == alice_id)) == 0
    assert db_session.scalar(select(func.count(VoteHistory.id)).where(VoteHistory.member_id == alice_id)) == 0
    assert db_session.scalar(select(func.count(PollAnswer.id)).where(PollAnswer.member_id == alice_id)) == 0

    # Bob still exists
    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == bob_id)) == 1

    # Alice's token gets 401
    me_res = client.get("/api/v1/me", headers={"X-Member-Token": alice_tok})
    assert me_res.status_code == 401


def test_delete_member_error_cases(client, db_session):
    """Member delete rejects missing member, wrong group, and unauthorized requests."""
    d1 = _setup_group_with_data(db_session, "Group 1")
    d2 = _setup_group_with_data(db_session, "Group 2")

    # 404 if member doesn't exist
    res = client.delete(
        f"/api/v1/groups/{d1['group'].id}/members/{uuid.uuid4()}",
        headers={"X-Admin-Token": d1["admin_token"]},
    )
    assert res.status_code == 404

    # 403 if member belongs to a different group (using d1 admin token to delete d2 member)
    res = client.delete(
        f"/api/v1/groups/{d1['group'].id}/members/{d2['members'][0].id}",
        headers={"X-Admin-Token": d1["admin_token"]},
    )
    assert res.status_code == 403
    assert res.json()["error"]["code"] == "forbidden"

    # 401 if missing token
    res = client.delete(f"/api/v1/groups/{d1['group'].id}/members/{d1['members'][0].id}")
    assert res.status_code == 401

    # 403 if token belongs to another group
    res = client.delete(
        f"/api/v1/groups/{d1['group'].id}/members/{d1['members'][0].id}",
        headers={"X-Admin-Token": d2["admin_token"]},
    )
    assert res.status_code == 403

    # Member token cannot delete
    res = client.delete(
        f"/api/v1/groups/{d1['group'].id}/members/{d1['members'][0].id}",
        headers={"X-Admin-Token": d1["alice_token"]},
    )
    assert res.status_code == 401


# ---------------------------------------------------------------------------
# 2. Bulk Member Deletion (POST /api/v1/groups/{group_id}/members/delete)
# ---------------------------------------------------------------------------

def test_bulk_delete_members_success(client, db_session):
    """Bulk delete successfully removes multiple members and their cascade data."""
    data = _setup_group_with_data(db_session, "Bulk Delete Group")
    group_id = data["group"].id
    admin_tok = data["admin_token"]
    m1_id = data["members"][0].id
    m2_id = data["members"][1].id
    m3_id = data["members"][2].id

    res = client.post(
        f"/api/v1/groups/{group_id}/members/delete",
        headers={"X-Admin-Token": admin_tok},
        json={"member_ids": [str(m1_id), str(m2_id)]},
    )
    assert res.status_code == 200
    assert res.json() == {"deleted": 2}

    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == m1_id)) == 0
    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == m2_id)) == 0
    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == m3_id)) == 1


def test_bulk_delete_all_or_nothing(client, db_session):
    """Bulk delete fails with 422 if any member does not belong to group, deleting nothing."""
    d1 = _setup_group_with_data(db_session, "Bulk Group 1")
    d2 = _setup_group_with_data(db_session, "Bulk Group 2")

    foreign_id = d2["members"][0].id
    valid_id = d1["members"][0].id

    res = client.post(
        f"/api/v1/groups/{d1['group'].id}/members/delete",
        headers={"X-Admin-Token": d1["admin_token"]},
        json={"member_ids": [str(valid_id), str(foreign_id)]},
    )
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "validation_error"

    # Nothing deleted
    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == valid_id)) == 1
    assert db_session.scalar(select(func.count(Member.id)).where(Member.id == foreign_id)) == 1


def test_bulk_delete_empty_and_duplicate_validation(client, db_session):
    """Bulk delete rejects empty list and duplicate IDs."""
    data = _setup_group_with_data(db_session, "Bulk Valid Group")
    group_id = data["group"].id
    admin_tok = data["admin_token"]
    mid = data["members"][0].id

    # Empty list
    res = client.post(
        f"/api/v1/groups/{group_id}/members/delete",
        headers={"X-Admin-Token": admin_tok},
        json={"member_ids": []},
    )
    assert res.status_code == 422

    # Duplicate IDs
    res = client.post(
        f"/api/v1/groups/{group_id}/members/delete",
        headers={"X-Admin-Token": admin_tok},
        json={"member_ids": [str(mid), str(mid)]},
    )
    assert res.status_code == 422
    assert "duplicate" in res.json()["error"]["message"].lower()


# ---------------------------------------------------------------------------
# 3. Poll Deletion (DELETE /api/v1/polls/{poll_id})
# ---------------------------------------------------------------------------

def test_delete_poll_cascades_options_votes_answers(client, db_session):
    """Deleting a poll removes everything for that poll and leaves the group and other polls untouched."""
    data = _setup_group_with_data(db_session, "Poll Delete Group")
    group_id = data["group"].id
    admin_tok = data["admin_token"]
    poll_id = data["poll"].id

    # Create a second poll
    second_poll = create_poll(
        db_session,
        group_id=group_id,
        name="Second Poll",
        allow_multiple=False,
        options=[PollOptionInput(label="A", role="target"), PollOptionInput(label="B", role="not_yet")],
    )
    db_session.commit()
    second_poll_id = second_poll.id

    # Verify poll has options, votes, answers, fields
    assert db_session.scalar(select(func.count(PollOption.id)).where(PollOption.poll_id == poll_id)) == 2
    assert db_session.scalar(select(func.count(Vote.id)).where(Vote.poll_id == poll_id)) == 1
    assert db_session.scalar(select(func.count(VoteHistory.id)).where(VoteHistory.poll_id == poll_id)) == 1
    assert db_session.scalar(select(func.count(PollField.id)).where(PollField.poll_id == poll_id)) == 1
    assert db_session.scalar(select(func.count(PollIncludedField.field_id)).where(PollIncludedField.poll_id == poll_id)) == 2
    assert db_session.scalar(select(func.count(PollAnswer.id)).where(PollAnswer.poll_id == poll_id)) == 1

    # Delete poll
    res = client.delete(
        f"/api/v1/polls/{poll_id}",
        headers={"X-Admin-Token": admin_tok},
    )
    assert res.status_code == 204

    # Poll and its children are gone
    assert db_session.scalar(select(func.count(Poll.id)).where(Poll.id == poll_id)) == 0
    assert db_session.scalar(select(func.count(PollOption.id)).where(PollOption.poll_id == poll_id)) == 0
    assert db_session.scalar(select(func.count(Vote.id)).where(Vote.poll_id == poll_id)) == 0
    assert db_session.scalar(select(func.count(VoteHistory.id)).where(VoteHistory.poll_id == poll_id)) == 0
    assert db_session.scalar(select(func.count(PollField.id)).where(PollField.poll_id == poll_id)) == 0
    assert db_session.scalar(select(func.count(PollIncludedField.field_id)).where(PollIncludedField.poll_id == poll_id)) == 0
    assert db_session.scalar(select(func.count(PollAnswer.id)).where(PollAnswer.poll_id == poll_id)) == 0

    # Second poll and group still exist
    assert db_session.scalar(select(func.count(Poll.id)).where(Poll.id == second_poll_id)) == 1
    assert db_session.scalar(select(func.count(Group.id)).where(Group.id == group_id)) == 1


def test_delete_poll_concurrent_waiting_vote_gets_404(db_session):
    """A vote transaction waiting on the poll lock gets clean PollNotFoundError (404) when poll is deleted."""
    session_factory = sessionmaker(bind=db_session.bind, autocommit=False, autoflush=False)

    data = _setup_group_with_data(db_session, "Lock Poll Group")
    poll_id = data["poll"].id
    member_id = data["members"][1].id  # Bob
    opt_id = data["poll"].options[0].id

    vote_result = {}
    lock_acquired = threading.Event()

    # Thread 1: Acquire lock on poll, signal lock_acquired, sleep, then delete poll
    def deleter():
        s = session_factory()
        try:
            poll = s.scalar(
                select(Poll).where(Poll.id == poll_id).with_for_update().execution_options(populate_existing=True)
            )
            lock_acquired.set()
            # Wait for voter to start waiting on this lock
            time.sleep(0.2)
            if poll:
                s.delete(poll)
                s.commit()
        finally:
            s.close()

    # Thread 2: Try to cast vote while deleter holds the lock
    def voter():
        s = session_factory()
        try:
            lock_acquired.wait(timeout=2.0)
            cast_vote(s, poll_id, member_id, opt_id)
            vote_result["status"] = "success"
        except PollNotFoundError:
            vote_result["status"] = "404"
        except Exception as e:
            vote_result["status"] = f"error: {type(e).__name__}"
        finally:
            s.close()

    t_del = threading.Thread(target=deleter)
    t_vot = threading.Thread(target=voter)

    t_del.start()
    t_vot.start()

    t_del.join(timeout=5.0)
    t_vot.join(timeout=5.0)

    # The voter MUST get 404 (PollNotFoundError), not server error
    assert vote_result.get("status") == "404"


# ---------------------------------------------------------------------------
# 4. Group Deletion (DELETE /api/v1/groups/{group_id})
# ---------------------------------------------------------------------------

def test_delete_group_confirmation_header(client, db_session):
    """Group delete requires exact name confirmation header (handles trim and rejects mismatches)."""
    data = _setup_group_with_data(db_session, "Exact Group Name")
    group_id = data["group"].id
    group_name = data["group"].name
    join_code = data["group"].join_code
    admin_tok = data["admin_token"]

    # Missing header -> 422 with message "Type the group name to confirm."
    res = client.delete(
        f"/api/v1/groups/{group_id}",
        headers={"X-Admin-Token": admin_tok},
    )
    assert res.status_code == 422
    assert res.json()["error"]["code"] == "validation_error"
    assert res.json()["error"]["message"] == "Type the group name to confirm."

    # Wrong name -> 422
    res = client.delete(
        f"/api/v1/groups/{group_id}",
        headers={"X-Admin-Token": admin_tok, "X-Confirm-Group-Name": "Wrong Name"},
    )
    assert res.status_code == 422
    assert res.json()["error"]["message"] == "Type the group name to confirm."

    # Extra spaces around exact name -> passes (trimmed)
    res = client.delete(
        f"/api/v1/groups/{group_id}",
        headers={"X-Admin-Token": admin_tok, "X-Confirm-Group-Name": f"   {group_name}   "},
    )
    assert res.status_code == 204

    # Group is deleted
    assert db_session.scalar(select(func.count(Group.id)).where(Group.id == group_id)) == 0

    # Join code returns 404
    join_res = client.get(f"/api/v1/join/{join_code}")
    assert join_res.status_code == 404

    # Admin call returns 404
    adm_res = client.get(f"/api/v1/groups/{group_id}", headers={"X-Admin-Token": admin_tok})
    assert adm_res.status_code == 404


# ---------------------------------------------------------------------------
# 5. Isolation: Deleting Group 1 Leaves All Tables in Group 2 Untouched
# ---------------------------------------------------------------------------

def test_delete_never_touches_other_groups(client, db_session):
    """Proves with exact counts that Group 2 data in EVERY table is completely unchanged."""
    d1 = _setup_group_with_data(db_session, "Group To Delete")
    d2 = _setup_group_with_data(db_session, "Group To Keep")

    g1_id = d1["group"].id
    g2_id = d2["group"].id

    g2_poll_subq = select(Poll.id).where(Poll.group_id == g2_id)

    # Capture Group 2 counts across all tables before deletion
    counts_before = {
        "groups": db_session.scalar(select(func.count(Group.id)).where(Group.id == g2_id)),
        "members": db_session.scalar(select(func.count(Member.id)).where(Member.group_id == g2_id)),
        "fields": db_session.scalar(select(func.count(GroupField.id)).where(GroupField.group_id == g2_id)),
        "polls": db_session.scalar(select(func.count(Poll.id)).where(Poll.group_id == g2_id)),
        "poll_options": db_session.scalar(
            select(func.count(PollOption.id)).where(PollOption.poll_id.in_(g2_poll_subq))
        ),
        "votes": db_session.scalar(
            select(func.count(Vote.id)).where(Vote.poll_id.in_(g2_poll_subq))
        ),
        "vote_history": db_session.scalar(
            select(func.count(VoteHistory.id)).where(VoteHistory.poll_id.in_(g2_poll_subq))
        ),
        "poll_fields": db_session.scalar(
            select(func.count(PollField.id)).where(PollField.poll_id.in_(g2_poll_subq))
        ),
        "poll_included_fields": db_session.scalar(
            select(func.count(PollIncludedField.field_id)).where(PollIncludedField.poll_id.in_(g2_poll_subq))
        ),
        "poll_answers": db_session.scalar(
            select(func.count(PollAnswer.id)).where(PollAnswer.poll_id.in_(g2_poll_subq))
        ),
    }

    # Verify counts are non-zero
    for table_name, count in counts_before.items():
        assert count > 0, f"Expected {table_name} for Group 2 to be > 0"

    # Delete Group 1
    res = client.delete(
        f"/api/v1/groups/{g1_id}",
        headers={"X-Admin-Token": d1["admin_token"], "X-Confirm-Group-Name": d1["group"].name},
    )
    assert res.status_code == 204

    # Group 1 rows are completely gone
    assert db_session.scalar(select(func.count(Group.id)).where(Group.id == g1_id)) == 0
    assert db_session.scalar(select(func.count(Member.id)).where(Member.group_id == g1_id)) == 0
    assert db_session.scalar(select(func.count(GroupField.id)).where(GroupField.group_id == g1_id)) == 0
    assert db_session.scalar(select(func.count(Poll.id)).where(Poll.group_id == g1_id)) == 0

    # Group 2 counts are EXACTLY identical across every table
    counts_after = {
        "groups": db_session.scalar(select(func.count(Group.id)).where(Group.id == g2_id)),
        "members": db_session.scalar(select(func.count(Member.id)).where(Member.group_id == g2_id)),
        "fields": db_session.scalar(select(func.count(GroupField.id)).where(GroupField.group_id == g2_id)),
        "polls": db_session.scalar(select(func.count(Poll.id)).where(Poll.group_id == g2_id)),
        "poll_options": db_session.scalar(
            select(func.count(PollOption.id)).where(PollOption.poll_id.in_(g2_poll_subq))
        ),
        "votes": db_session.scalar(
            select(func.count(Vote.id)).where(Vote.poll_id.in_(g2_poll_subq))
        ),
        "vote_history": db_session.scalar(
            select(func.count(VoteHistory.id)).where(VoteHistory.poll_id.in_(g2_poll_subq))
        ),
        "poll_fields": db_session.scalar(
            select(func.count(PollField.id)).where(PollField.poll_id.in_(g2_poll_subq))
        ),
        "poll_included_fields": db_session.scalar(
            select(func.count(PollIncludedField.field_id)).where(PollIncludedField.poll_id.in_(g2_poll_subq))
        ),
        "poll_answers": db_session.scalar(
            select(func.count(PollAnswer.id)).where(PollAnswer.poll_id.in_(g2_poll_subq))
        ),
    }

    assert counts_before == counts_after


# ---------------------------------------------------------------------------
# 6. Roster Export (GET /api/v1/groups/{group_id}/members/export)
# ---------------------------------------------------------------------------

def test_roster_export_csv_and_xlsx_with_formula_protection(client, db_session):
    """Roster export produces correct CSV and Excel files with headers, order, and formula protection."""
    data = _setup_group_with_data(db_session, "Export Group 101")
    group_id = data["group"].id
    admin_tok = data["admin_token"]

    # 1. Test CSV export
    res_csv = client.get(
        f"/api/v1/groups/{group_id}/members/export?format=csv",
        headers={"X-Admin-Token": admin_tok},
    )
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    assert 'filename="roster-Export_Group_101.csv"' in res_csv.headers["content-disposition"]

    csv_text = res_csv.content.decode("utf-8-sig")
    lines = [line.strip() for line in csv_text.splitlines() if line.strip()]
    header = lines[0].split(",")
    # Columns: Name, Department, RegisterNo, Status, Active
    assert header == ["Name", "Department", "RegisterNo", "Status", "Active"]

    # Check Charlie's row has formula protection: starts with '=Charlie' -> should be "'=Charlie"
    # and Department has '+Lead' -> should be "'+Lead"
    charlie_line = [l for l in lines if "Charlie" in l][0]
    assert "'=Charlie" in charlie_line
    assert "'+Lead" in charlie_line
    assert "'-R103" in charlie_line

    # 2. Test XLSX export
    res_xlsx = client.get(
        f"/api/v1/groups/{group_id}/members/export?format=xlsx",
        headers={"X-Admin-Token": admin_tok},
    )
    assert res_xlsx.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_xlsx.headers["content-type"]
    assert 'filename="roster-Export_Group_101.xlsx"' in res_xlsx.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    ws = wb.active
    assert ws.title == "Roster"
    header_vals = [ws.cell(1, c).value for c in range(1, 6)]
    assert header_vals == ["Name", "Department", "RegisterNo", "Status", "Active"]

    # Find Charlie's row in Excel
    charlie_row = None
    for r in range(2, ws.max_row + 1):
        if "Charlie" in str(ws.cell(r, 1).value):
            charlie_row = r
            break
    assert charlie_row is not None
    # Check value starts with single quote for formula protection
    assert ws.cell(charlie_row, 1).value == "'=Charlie"
    assert ws.cell(charlie_row, 2).value == "'+Lead"
    assert ws.cell(charlie_row, 3).value == "'-R103"


def test_roster_export_auth(client, db_session):
    """Roster export requires admin token; rejects missing or invalid token."""
    data = _setup_group_with_data(db_session, "Auth Export Group")
    group_id = data["group"].id

    # Missing token -> 401
    res = client.get(f"/api/v1/groups/{group_id}/members/export?format=csv")
    assert res.status_code == 401

    # Member token -> 401
    res = client.get(
        f"/api/v1/groups/{group_id}/members/export?format=csv",
        headers={"X-Admin-Token": data["alice_token"]},
    )
    assert res.status_code == 401
