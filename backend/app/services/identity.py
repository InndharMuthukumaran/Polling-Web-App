"""Identity service layer implementing group creation, admin/member tokens, and claim management."""

import csv
from datetime import datetime, timezone
import io
import re
from typing import Any
import uuid

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.errors import (
    ClaimConflictError,
    ClaimNotApprovedError,
    ClaimStateError,
    GroupNotFoundError,
    InvalidTokenError,
    MemberGroupMismatchError,
    MemberInactiveError,
    MemberNotFoundError,
    PermissionDeniedError,
    PollValidationError,
)
from app.limits import MAX_MEMBERS_PER_REQUEST
from app.models import Group, GroupField, Member
from app.security import (
    generate_join_code,
    generate_token,
    hash_token,
    verify_token_hash,
)
from app.services.fields import (
    normalize_identifier,
    validate_member_values_for_create,
    validate_member_values_for_update,
)
from app.services.polls import sanitize_formula_injection


def ensure_utc(dt: datetime | None) -> datetime | None:
    """Ensure a datetime is timezone-aware UTC."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def resolve_uuid(val: uuid.UUID | str) -> uuid.UUID:
    """Resolve a UUID from either UUID or string representation."""
    if isinstance(val, uuid.UUID):
        return val
    return uuid.UUID(str(val))


def create_group_with_admin(
    session: Session,
    name: str,
    now: datetime | None = None,
) -> tuple[Group, str]:
    """Create a new group with a join code and admin token.

    Returns the created Group and the plain, unhashed admin token.
    """
    if not name or not name.strip():
        raise PollValidationError("Group name cannot be empty.")

    current_time = ensure_utc(now) or datetime.now(timezone.utc)
    admin_token = generate_token()
    join_code = generate_join_code()

    group = Group(
        id=uuid.uuid4(),
        name=name.strip(),
        join_code=join_code,
        admin_token_hash=hash_token(admin_token),
        require_claim_approval=False,
        created_at=current_time,
    )
    session.add(group)
    session.commit()
    session.refresh(group)
    return group, admin_token


def authenticate_admin(
    session: Session,
    group_id: uuid.UUID | str,
    admin_token: str,
) -> Group:
    """Authenticate an admin token against a group.

    Raises GroupNotFoundError if group does not exist.
    Raises InvalidTokenError if token is empty, wrong, or group lacks an admin token.
    """
    gid = resolve_uuid(group_id)
    group = session.get(Group, gid)
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    if not admin_token or not isinstance(admin_token, str) or not admin_token.strip():
        raise InvalidTokenError("Admin token cannot be empty.")

    if not group.admin_token_hash:
        raise InvalidTokenError("Group has no admin token configured.")

    if not verify_token_hash(admin_token.strip(), group.admin_token_hash):
        raise InvalidTokenError("Invalid admin token.")

    return group


def update_group_settings(
    session: Session,
    group_id: uuid.UUID | str,
    require_claim_approval: bool | None = None,
    allow_name_list: bool | None = None,
) -> Group:
    """Update settings for a group."""
    gid = resolve_uuid(group_id)
    group = session.get(Group, gid)
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    if require_claim_approval is not None:
        group.require_claim_approval = bool(require_claim_approval)
    if allow_name_list is not None:
        group.allow_name_list = bool(allow_name_list)

    session.commit()
    session.refresh(group)
    return group


def add_members_bulk(
    session: Session,
    group_id: uuid.UUID | str,
    display_names: list[str] | list[dict[str, Any]] | None = None,
    members_data: list[dict[str, Any]] | None = None,
) -> list[Member]:
    """Bulk-add members to a group. All-or-nothing validation with row-level details."""
    gid = resolve_uuid(group_id)

    # Concurrency lock on group row
    group = session.scalar(
        select(Group).where(Group.id == gid).with_for_update()
    )
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    raw_items: list[dict[str, Any]] = []
    if members_data is not None:
        raw_items = members_data
    elif display_names is not None:
        for item in display_names:
            if isinstance(item, str):
                raw_items.append({"display_name": item, "values": {}})
            elif isinstance(item, dict):
                raw_items.append(item)
            else:
                raw_items.append({"display_name": item, "values": {}})

    if len(raw_items) > MAX_MEMBERS_PER_REQUEST:
        raise PollValidationError(f"Cannot add more than {MAX_MEMBERS_PER_REQUEST} members in a single request.")
    if len(raw_items) == 0:
        raise PollValidationError("Members list cannot be empty.")

    fields = list(
        session.scalars(
            select(GroupField)
            .where(GroupField.group_id == gid)
            .order_by(GroupField.position)
        ).all()
    )
    id_field = next((f for f in fields if f.is_identifier), None)

    # Query existing members for uniqueness checks
    existing_members = list(
        session.scalars(select(Member).where(Member.group_id == gid)).all()
    )
    existing_display_names: set[str] = {m.display_name for m in existing_members}
    existing_identifiers: set[str] = {
        m.identifier_value for m in existing_members if m.identifier_value is not None
    }

    details: list[dict[str, Any]] = []
    seen_display_names: dict[str, int] = {}
    seen_identifiers: dict[str, int] = {}
    validated_rows: list[tuple[str, dict[str, Any], str | None]] = []

    for idx, item in enumerate(raw_items):
        row = idx + 1
        name_val = item.get("display_name")
        if not isinstance(name_val, str) or not name_val.strip():
            details.append({
                "row": row,
                "field": "display_name",
                "message": "Member display_name cannot be blank.",
            })
            clean_name = ""
        else:
            clean_name = name_val.strip()

        # Check values
        raw_vals = item.get("values")
        cleaned_vals: dict[str, Any] = {}
        identifier_val: str | None = None
        try:
            cleaned_vals, identifier_val = validate_member_values_for_create(
                fields, raw_vals, row=row
            )
        except PollValidationError as exc:
            if exc.details:
                details.extend(exc.details)
            else:
                details.append({
                    "row": row,
                    "field": "values",
                    "message": exc.message,
                })

        # Uniqueness checks
        if id_field is not None:
            # Group has an identifier field: duplicate display names allowed, identifier must be unique
            if identifier_val is not None:
                if identifier_val in seen_identifiers:
                    details.append({
                        "row": row,
                        "field": id_field.name,
                        "message": "Duplicate identifier value in batch.",
                    })
                else:
                    seen_identifiers[identifier_val] = row

                if identifier_val in existing_identifiers:
                    details.append({
                        "row": row,
                        "field": id_field.name,
                        "message": "Member with identifier already exists in group.",
                    })
        else:
            # Group has NO identifier field: display_name must be unique
            if clean_name:
                if clean_name in seen_display_names:
                    details.append({
                        "row": row,
                        "field": "display_name",
                        "message": f"Duplicate display_name in batch: '{clean_name}'.",
                    })
                else:
                    seen_display_names[clean_name] = row

                if clean_name in existing_display_names:
                    details.append({
                        "row": row,
                        "field": "display_name",
                        "message": f"Member with name '{clean_name}' already exists in group.",
                    })

        validated_rows.append((clean_name, cleaned_vals, identifier_val))

    if details:
        primary_msg = details[0]["message"] if len(details) == 1 else "Validation failed for one or more members."
        raise PollValidationError(
            primary_msg,
            details=details,
        )

    current_time = datetime.now(timezone.utc)
    new_members: list[Member] = []
    for clean_name, cleaned_vals, identifier_val in validated_rows:
        m = Member(
            id=uuid.uuid4(),
            group_id=gid,
            display_name=clean_name,
            is_active=True,
            claim_status="unclaimed",
            member_token_hash=None,
            claimed_at=None,
            field_values=cleaned_vals,
            identifier_value=identifier_val,
            created_at=current_time,
        )
        if id_field is not None:
            raw_id = cleaned_vals.get(id_field.key)
            m.identifier = str(raw_id) if raw_id is not None else None
        else:
            m.identifier = None
        session.add(m)
        new_members.append(m)

    session.commit()
    for m in new_members:
        session.refresh(m)
        if id_field is not None:
            raw_id = (m.field_values or {}).get(id_field.key)
            m.identifier = str(raw_id) if raw_id is not None else None
        else:
            m.identifier = None
    return new_members


def update_member(
    session: Session,
    group_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
    display_name: str | None = None,
    is_active: bool | None = None,
    values: dict[str, Any] | None = None,
) -> Member:
    """Update member display name, active status, or custom field values."""
    gid = resolve_uuid(group_id)
    mid = resolve_uuid(member_id)

    # Concurrency lock
    group = session.scalar(
        select(Group).where(Group.id == gid).with_for_update()
    )
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    member = session.get(Member, mid)
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != gid:
        raise PermissionDeniedError(f"Member {mid} does not belong to group {gid}.")

    fields = list(
        session.scalars(
            select(GroupField)
            .where(GroupField.group_id == gid)
            .order_by(GroupField.position)
        ).all()
    )
    id_field = next((f for f in fields if f.is_identifier), None)

    if display_name is not None:
        if not isinstance(display_name, str) or not display_name.strip():
            raise PollValidationError("display_name cannot be blank.")
        trimmed = display_name.strip()
        if trimmed != member.display_name:
            if id_field is None:
                # Uniqueness required only when no identifier field exists
                existing = session.scalar(
                    select(Member).where(
                        Member.group_id == gid,
                        Member.display_name == trimmed,
                        Member.id != mid,
                    )
                )
                if existing:
                    raise PollValidationError(f"Display name '{trimmed}' already in use.")
            member.display_name = trimmed

    if is_active is not None:
        member.is_active = bool(is_active)

    if values is not None:
        updated_values, new_id_val = validate_member_values_for_update(
            fields, member.field_values or {}, values
        )
        if id_field is not None and new_id_val is not None:
            # Check identifier uniqueness among other members
            existing = session.scalar(
                select(Member).where(
                    Member.group_id == gid,
                    Member.identifier_value == new_id_val,
                    Member.id != mid,
                )
            )
            if existing:
                raise PollValidationError(
                    f"Identifier value '{new_id_val}' already in use in group."
                )
            member.identifier_value = new_id_val

        # Assign a NEW dict so SQLAlchemy sees the mutation
        member.field_values = {**updated_values}

    session.commit()
    session.refresh(member)
    if id_field is not None:
        raw_id = (member.field_values or {}).get(id_field.key)
        member.identifier = str(raw_id) if raw_id is not None else None
    else:
        member.identifier = None
    return member


def list_group_members(
    session: Session,
    group_id: uuid.UUID | str,
) -> list[Member]:
    """List all members of a group ordered by display name."""
    gid = resolve_uuid(group_id)
    group = session.get(Group, gid)
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    fields = list(
        session.scalars(
            select(GroupField).where(GroupField.group_id == gid)
        ).all()
    )
    id_field = next((f for f in fields if f.is_identifier), None)

    members = list(
        session.scalars(
            select(Member)
            .where(Member.group_id == gid)
            .order_by(Member.display_name)
        ).all()
    )
    for m in members:
        if id_field and m.field_values:
            raw_id = m.field_values.get(id_field.key)
            m.identifier = str(raw_id) if raw_id is not None else None
        else:
            m.identifier = None
    return members


def get_group_by_join_code(
    session: Session,
    join_code: str,
) -> Group:
    """Retrieve group by its unguessable join code."""
    if not join_code or not isinstance(join_code, str) or not join_code.strip():
        raise GroupNotFoundError("Join code cannot be empty.")

    group = session.scalar(
        select(Group).where(Group.join_code == join_code.strip())
    )
    if not group:
        raise GroupNotFoundError(f"No group found for join code: '{join_code}'.")
    return group


def claim_member(
    session: Session,
    join_code: str,
    member_id: uuid.UUID | str,
    now: datetime | None = None,
) -> tuple[Member, str]:
    """Claim a member name using the group join code.

    Locks member row with FOR UPDATE. Returns (member, plain_member_token).
    """
    group = get_group_by_join_code(session, join_code)
    mid = resolve_uuid(member_id)

    # Lock the member row
    member = session.scalar(
        select(Member).where(Member.id == mid).with_for_update()
    )
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != group.id:
        raise PermissionDeniedError(
            f"Member {mid} does not belong to group with join code '{join_code}'."
        )

    if not member.is_active:
        raise MemberInactiveError(f"Member {mid} is inactive.")

    if member.claim_status != "unclaimed":
        raise ClaimConflictError(
            f"Member {mid} is already claimed with status '{member.claim_status}'."
        )

    token = generate_token()
    claim_time = ensure_utc(now) or datetime.now(timezone.utc)

    member.member_token_hash = hash_token(token)
    member.claimed_at = claim_time
    member.claim_status = "pending" if group.require_claim_approval else "approved"

    session.commit()
    session.refresh(member)
    return member, token


def approve_claim(
    session: Session,
    group_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
) -> Member:
    """Approve a pending member claim."""
    gid = resolve_uuid(group_id)
    mid = resolve_uuid(member_id)

    member = session.scalar(
        select(Member).where(Member.id == mid).with_for_update()
    )
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != gid:
        raise PermissionDeniedError(f"Member {mid} does not belong to group {gid}.")

    if member.claim_status == "unclaimed":
        raise ClaimStateError(f"Cannot approve claim for unclaimed member {mid}.")

    if member.claim_status == "pending":
        member.claim_status = "approved"
        session.commit()
        session.refresh(member)

    # Already approved is a no-op
    return member


def reset_claim(
    session: Session,
    group_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
) -> Member:
    """Reset a member claim to unclaimed, revoking any active token.

    Existing votes and history remain intact.
    """
    gid = resolve_uuid(group_id)
    mid = resolve_uuid(member_id)

    member = session.scalar(
        select(Member).where(Member.id == mid).with_for_update()
    )
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != gid:
        raise PermissionDeniedError(f"Member {mid} does not belong to group {gid}.")

    member.claim_status = "unclaimed"
    member.member_token_hash = None
    member.claimed_at = None

    session.commit()
    session.refresh(member)
    return member


def authenticate_member(
    session: Session,
    member_token: str,
) -> Member:
    """Authenticate a member by plain member token.

    Does not verify approval status.
    """
    if not member_token or not isinstance(member_token, str) or not member_token.strip():
        raise InvalidTokenError("Member token cannot be empty.")

    token_hash = hash_token(member_token.strip())
    member = session.scalar(
        select(Member).where(Member.member_token_hash == token_hash)
    )
    if not member:
        raise InvalidTokenError("Invalid member token.")

    return member


def require_approved(member: Member) -> None:
    """Raise ClaimNotApprovedError if member claim_status is not 'approved'."""
    if member.claim_status != "approved":
        raise ClaimNotApprovedError(
            f"Member '{member.display_name}' claim is not approved (status: {member.claim_status})."
        )


def release_own_claim(
    session: Session,
    member: Member,
) -> Member:
    """Release a member's own claim, returning claim_status to unclaimed and clearing token hash.

    Reuses reset_claim logic so votes and history are untouched.
    """
    return reset_claim(session, member.group_id, member.id)


def mask_identifier(val: str | None) -> str | None:
    """Mask an identifier for public display.

    Shows only the last 3 characters prefixed by '•••' for values longer than 3
    characters (e.g. '•••001'). Values of 3 or fewer characters are fully masked
    with '•' so that the full identifier is never exposed on a public endpoint.
    """
    if val is None:
        return None
    val_str = str(val)
    if not val_str:
        return ""
    if len(val_str) > 3:
        return f"•••{val_str[-3:]}"
    return "•" * len(val_str)


def lookup_member_by_identifier(
    session: Session,
    join_code: str,
    identifier: str,
) -> Member:
    """Lookup an active member by identifier within a group.

    Normalizes input using the same logic as R1a (trim, collapse whitespace, casefold).
    If the group has no identifier field, or there is no match, or the member is inactive,
    or the join code is invalid, raises MemberNotFoundError with a generic message.
    """
    if not join_code or not isinstance(join_code, str) or not join_code.strip():
        raise MemberNotFoundError("No member matches that identifier.")

    group = session.scalar(
        select(Group).where(Group.join_code == join_code.strip())
    )
    if not group:
        raise MemberNotFoundError("No member matches that identifier.")

    id_field = session.scalar(
        select(GroupField).where(
            GroupField.group_id == group.id,
            GroupField.is_identifier.is_(True),
        )
    )
    if not id_field:
        raise MemberNotFoundError("No member matches that identifier.")

    if not identifier or not isinstance(identifier, str):
        raise MemberNotFoundError("No member matches that identifier.")

    normalized = normalize_identifier(identifier)
    if not normalized:
        raise MemberNotFoundError("No member matches that identifier.")

    member = session.scalar(
        select(Member).where(
            Member.group_id == group.id,
            Member.identifier_value == normalized,
            Member.is_active.is_(True),
        )
    )
    if not member:
        raise MemberNotFoundError("No member matches that identifier.")

    return member


def delete_member(
    session: Session,
    group_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
) -> None:
    """Permanently delete a member and all their votes, history, and poll answers.

    Locks the group row first, checks that the member exists and belongs to the group,
    deletes the member, and commits.
    """
    gid = resolve_uuid(group_id)
    mid = resolve_uuid(member_id)

    # Lock the group row first
    group = session.scalar(
        select(Group).where(Group.id == gid).with_for_update()
    )
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    member = session.get(Member, mid)
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != group.id:
        raise MemberGroupMismatchError("Member belongs to a different group.")

    session.delete(member)
    session.commit()


def delete_members_bulk(
    session: Session,
    group_id: uuid.UUID | str,
    member_ids: list[uuid.UUID | str],
) -> int:
    """Bulk delete members within a group (all-or-nothing).

    Validates that:
    - member_ids has at least 1 item
    - member_ids does not exceed MAX_MEMBERS_PER_REQUEST
    - member_ids contains no duplicates
    - every member_id exists and belongs to this group

    If any check fails, raises PollValidationError (422) and deletes nothing.
    Returns the count of deleted members.
    """
    if not member_ids:
        raise PollValidationError("At least one member ID must be provided.")

    if len(member_ids) > MAX_MEMBERS_PER_REQUEST:
        raise PollValidationError(
            f"Cannot delete more than {MAX_MEMBERS_PER_REQUEST} members at once."
        )

    uids = [resolve_uuid(mid) for mid in member_ids]
    if len(uids) != len(set(uids)):
        raise PollValidationError("Duplicate member IDs are not allowed.")

    gid = resolve_uuid(group_id)
    # Lock the group row first
    group = session.scalar(
        select(Group).where(Group.id == gid).with_for_update()
    )
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    members = list(
        session.scalars(
            select(Member).where(
                Member.group_id == group.id,
                Member.id.in_(uids),
            )
        ).all()
    )

    if len(members) != len(uids):
        raise PollValidationError(
            "One or more member IDs do not belong to this group or do not exist."
        )

    for m in members:
        session.delete(m)

    session.commit()
    return len(members)


def delete_group(
    session: Session,
    group_id: uuid.UUID | str,
) -> None:
    """Permanently delete a group and all its members, fields, polls, votes, and answers.

    Locks the group row with with_for_update() first, deletes it, and commits.
    Cascades remove all associated entities.
    """
    gid = resolve_uuid(group_id)
    group = session.scalar(
        select(Group).where(Group.id == gid).with_for_update()
    )
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    session.delete(group)
    session.commit()


def make_group_slug(name: str) -> str:
    """Generate safe filename slug from group name."""
    safe = re.sub(r'[^a-zA-Z0-9_\-]+', '_', name.strip())
    safe = safe.strip('_')
    return safe[:100] if safe else "group"


def get_roster_data(
    session: Session,
    group_id: uuid.UUID | str,
) -> tuple[Group, list[GroupField], list[Member]]:
    """Retrieve group, ordered fields, and members for roster export."""
    gid = resolve_uuid(group_id)
    group = session.get(Group, gid)
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    fields = list(
        session.scalars(
            select(GroupField)
            .where(GroupField.group_id == group.id)
            .order_by(GroupField.position)
        ).all()
    )

    members = list(
        session.scalars(
            select(Member)
            .where(Member.group_id == group.id)
            .order_by(Member.created_at.asc())
        ).all()
    )

    return group, fields, members


def export_roster_csv(session: Session, group_id: uuid.UUID | str) -> tuple[bytes, str]:
    """Generate CSV bytes (UTF-8 with BOM) and safe filename for group roster."""
    group, fields, members = get_roster_data(session, group_id)
    slug = make_group_slug(group.name)
    filename = f"roster-{slug}.csv"

    headers = ["Name"] + [f.name for f in fields] + ["Status", "Active"]

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([sanitize_formula_injection(h) for h in headers])

    status_map = {
        "unclaimed": "Not claimed",
        "pending": "Waiting for approval",
        "approved": "Claimed",
    }

    for m in members:
        status_label = status_map.get(m.claim_status, m.claim_status)
        active_label = "Yes" if m.is_active else "No"
        row_cells: list[Any] = [m.display_name]
        for f in fields:
            val = (m.field_values or {}).get(f.key, "")
            row_cells.append(val if val is not None else "")
        row_cells.append(status_label)
        row_cells.append(active_label)

        writer.writerow([sanitize_formula_injection(cell) for cell in row_cells])

    csv_bytes = output.getvalue().encode("utf-8-sig")
    return csv_bytes, filename


def export_roster_xlsx(session: Session, group_id: uuid.UUID | str) -> tuple[bytes, str]:
    """Generate Excel (.xlsx) bytes and safe filename for group roster."""
    group, fields, members = get_roster_data(session, group_id)
    slug = make_group_slug(group.name)
    filename = f"roster-{slug}.xlsx"

    headers = ["Name"] + [f.name for f in fields] + ["Status", "Active"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"

    for c_idx, h in enumerate(headers, start=1):
        sanitized = sanitize_formula_injection(h)
        cell = ws.cell(row=1, column=c_idx, value=str(sanitized))
        cell.data_type = "s"

    status_map = {
        "unclaimed": "Not claimed",
        "pending": "Waiting for approval",
        "approved": "Claimed",
    }

    for r_idx, m in enumerate(members, start=2):
        status_label = status_map.get(m.claim_status, m.claim_status)
        active_label = "Yes" if m.is_active else "No"
        row_cells: list[Any] = [m.display_name]
        for f in fields:
            val = (m.field_values or {}).get(f.key, "")
            row_cells.append(val if val is not None else "")
        row_cells.append(status_label)
        row_cells.append(active_label)

        for c_idx, val in enumerate(row_cells, start=1):
            sanitized = sanitize_formula_injection(val)
            cell = ws.cell(row=r_idx, column=c_idx)
            if isinstance(sanitized, (int, float)) and not isinstance(sanitized, bool):
                cell.value = sanitized
                cell.data_type = "n"
            else:
                cell.value = str(sanitized) if sanitized is not None else ""
                cell.data_type = "s"

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()
    return xlsx_bytes, filename

