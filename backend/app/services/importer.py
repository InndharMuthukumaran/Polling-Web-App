"""Spreadsheet importer service for bulk member onboarding.

Handles file validation, parsing (.xlsx and .csv), cell normalization,
mapping suggestions, template generation, and transactional all-or-nothing imports.
"""

from collections import Counter
import csv
import datetime
import io
import math
import re
from typing import Any, Sequence
import uuid
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.errors import GroupNotFoundError, PollValidationError
from app.models import Group, GroupField, Member
from app.services.fields import (
    create_group_field,
    is_blank,
    list_group_fields,
    normalize_identifier,
    validate_member_values_for_create,
)
from app.services.identity import add_members_bulk, resolve_uuid

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MAX_UNCOMPRESSED_SIZE = 50 * 1024 * 1024  # 50 MB
MAX_DATA_ROWS = 2000
MAX_COLUMNS = 60
MAX_ERROR_DETAILS = 100


def read_upload_capped(file: Any, max_bytes: int = MAX_FILE_SIZE) -> bytes:
    """Read at most max_bytes + 1 bytes from an uploaded file.

    Raises 422 PollValidationError without reading further if upload exceeds max_bytes.
    """
    content = file.file.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise PollValidationError("File size exceeds 5 MB limit.")
    return content


def validate_file_type_and_size(filename: str, content: bytes) -> str:
    """Validate file size, extension, and signature. Returns 'xlsx' or 'csv'."""
    if len(content) > MAX_FILE_SIZE:
        raise PollValidationError("File size exceeds 5 MB limit.")

    fname_lower = filename.lower()
    if fname_lower.endswith(".xls") or fname_lower.endswith(".xlsm"):
        raise PollValidationError("Please save the file as .xlsx or .csv.")

    if fname_lower.endswith(".xlsx"):
        file_type = "xlsx"
    elif fname_lower.endswith(".csv"):
        file_type = "csv"
    else:
        raise PollValidationError("Unsupported file type. Please upload a .xlsx or .csv file.")

    if file_type == "xlsx":
        if not content.startswith(b"PK"):
            raise PollValidationError("Invalid file content. Not a valid .xlsx file.")

        try:
            zf = zipfile.ZipFile(io.BytesIO(content))
            uncompressed_size = sum(info.file_size for info in zf.infolist())
            if uncompressed_size > MAX_UNCOMPRESSED_SIZE:
                raise PollValidationError("Uncompressed file size exceeds 50 MB limit.")
        except zipfile.BadZipFile:
            raise PollValidationError("Invalid file content. Corrupted or invalid .xlsx file.")

    return file_type


def convert_cell_value(val: Any, target_field_type: str | None = None) -> Any:
    """Convert raw spreadsheet cell value to expected type before validation."""
    if val is None:
        return ""

    if isinstance(val, bool):
        return "TRUE" if val else "FALSE"

    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime("%Y-%m-%d")

    if isinstance(val, float):
        if not math.isfinite(val):
            return str(val)
        if val.is_integer():
            if target_field_type == "number":
                return int(val)
            return str(int(val))
        if target_field_type == "number":
            return val
        return str(val)

    if isinstance(val, int):
        if target_field_type == "number":
            return val
        return str(val)

    if isinstance(val, str):
        return val.strip()

    return str(val).strip()


def parse_raw_spreadsheet(
    filename: str, content: bytes
) -> tuple[str | None, list[str], list[tuple[int, list[Any]]]]:
    """Parse spreadsheet content into (sheet_name, headers, list of (sheet_row_num, row_cells)).

    Skips fully blank leading rows. The header row is the first row with any non-blank cell.
    Subsequent fully blank rows are skipped.
    """
    file_type = validate_file_type_and_size(filename, content)

    raw_rows: list[tuple[int, list[Any]]] = []
    sheet_name: str | None = None

    if file_type == "xlsx":
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            sheet_name = ws.title
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                raw_rows.append((row_idx, list(row)))
            wb.close()
        except PollValidationError:
            raise
        except Exception as exc:
            raise PollValidationError(f"Unable to read Excel workbook: {exc}")
    else:
        # CSV parsing
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1252")
            except UnicodeDecodeError:
                raise PollValidationError("Unable to decode CSV file with UTF-8 or CP1252.")

        delimiter = ","
        try:
            sample = text[:4096]
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            if dialect.delimiter in (",", ";", "\t"):
                delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        for row_idx, row in enumerate(reader, start=1):
            raw_rows.append((row_idx, row))

    def is_blank_row(cells: list[Any]) -> bool:
        return all(c is None or str(c).strip() == "" for c in cells)

    # Find the header row (first non-blank row)
    header_idx = -1
    for i, (_, cells) in enumerate(raw_rows):
        if not is_blank_row(cells):
            header_idx = i
            break

    if header_idx == -1:
        raise PollValidationError("The file is empty or contains no headers.")

    header_sheet_row, header_cells = raw_rows[header_idx]

    # Find the last non-empty header column
    last_col = -1
    for c_idx in range(len(header_cells) - 1, -1, -1):
        c = header_cells[c_idx]
        if c is not None and str(c).strip() != "":
            last_col = c_idx
            break

    if last_col == -1:
        raise PollValidationError("The file contains no headers.")

    active_headers_raw = header_cells[: last_col + 1]
    if len(active_headers_raw) > MAX_COLUMNS:
        raise PollValidationError("File contains more than the maximum limit of 60 columns.")

    headers: list[str] = []
    seen_headers: dict[str, int] = {}
    for col_idx, h in enumerate(active_headers_raw, start=1):
        if h is None or str(h).strip() == "":
            raise PollValidationError(f"Blank header at column {col_idx}.")
        h_str = str(h).strip()
        h_cf = h_str.casefold()
        if h_cf in seen_headers:
            prev_col = seen_headers[h_cf]
            raise PollValidationError(
                f"Duplicate header '{h_str}' at column {col_idx} (already defined at column {prev_col})."
            )
        seen_headers[h_cf] = col_idx
        headers.append(h_str)

    # Process data rows
    data_rows: list[tuple[int, list[Any]]] = []
    for sheet_row_num, cells in raw_rows[header_idx + 1 :]:
        if is_blank_row(cells):
            continue
        padded = cells[: len(headers)] + [None] * max(0, len(headers) - len(cells))
        data_rows.append((sheet_row_num, padded))

    if len(data_rows) > MAX_DATA_ROWS:
        raise PollValidationError(
            f"File contains more than the maximum limit of {MAX_DATA_ROWS} data rows."
        )

    return sheet_name, headers, data_rows


def generate_suggested_mapping(
    headers: list[str], fields: list[GroupField]
) -> dict[str, str]:
    """Suggest column mapping to 'name', a field key, or 'skip'."""
    suggestions: dict[str, str] = {}
    field_by_name = {f.name.casefold(): f.key for f in fields}
    field_by_key = {f.key.casefold(): f.key for f in fields}

    assigned_targets: set[str] = set()

    for i, h in enumerate(headers):
        h_norm = " ".join(h.strip().split()).casefold()
        target = "skip"

        if h_norm in ("name", "full name", "student name") and "name" not in assigned_targets:
            target = "name"
            assigned_targets.add("name")
        elif h_norm in field_by_name and field_by_name[h_norm] not in assigned_targets:
            target = field_by_name[h_norm]
            assigned_targets.add(target)
        elif h_norm in field_by_key and field_by_key[h_norm] not in assigned_targets:
            target = field_by_key[h_norm]
            assigned_targets.add(target)

        suggestions[str(i)] = target

    return suggestions


def preview_import_data(
    filename: str,
    content: bytes,
    fields: list[GroupField],
) -> dict[str, Any]:
    """Generate preview data from file."""
    sheet_name, headers, data_rows = parse_raw_spreadsheet(filename, content)

    columns = [{"index": i, "header": h} for i, h in enumerate(headers)]
    sample_rows = [
        [convert_cell_value(cell) for cell in cells]
        for _, cells in data_rows[:10]
    ]

    suggested = generate_suggested_mapping(headers, fields)

    return {
        "filename": filename,
        "sheet": sheet_name,
        "columns": columns,
        "total_rows": len(data_rows),
        "sample_rows": sample_rows,
        "suggested_mapping": suggested,
    }


def generate_template_file(
    fields: list[GroupField],
    file_format: str,
) -> tuple[bytes, str, str]:
    """Generate a spreadsheet template (.xlsx or .csv) with headers and choice dropdowns."""
    sorted_fields = sorted(fields, key=lambda f: f.position)
    header_row = ["Name"] + [f.name for f in sorted_fields]

    if file_format.lower() == "csv":
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(header_row)
        return (
            out.getvalue().encode("utf-8"),
            "text/csv; charset=utf-8",
            "members-template.csv",
        )

    # Default to xlsx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    ws.append(header_row)

    # Add dropdown validation for choice fields
    for col_idx, f in enumerate(sorted_fields, start=2):
        if f.field_type == "choice" and f.choices:
            col_letter = get_column_letter(col_idx)
            escaped = ",".join(str(c).replace('"', '""') for c in f.choices)
            formula = f'"{escaped}"'
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col_letter}2:{col_letter}2001")

    bio = io.BytesIO()
    wb.save(bio)
    wb.close()
    return (
        bio.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "members-template.xlsx",
    )


def execute_spreadsheet_import(
    session: Session,
    group: Group,
    filename: str,
    content: bytes,
    mapping: dict[str, str],
    new_fields_data: list[dict[str, Any]] | None = None,
    dry_run: bool = False,
    on_duplicate: str = "reject",
) -> dict[str, Any]:
    """Execute bulk member import with all-or-nothing validation."""
    if on_duplicate not in ("reject", "skip"):
        raise PollValidationError("on_duplicate must be either 'reject' or 'skip'.")

    # 1. Parse spreadsheet
    sheet_name, headers, data_rows = parse_raw_spreadsheet(filename, content)
    num_cols = len(headers)

    # 2. Validate mapping structure
    # Mapping keys may be integers or strings representing column indices
    col_to_target: dict[int, str] = {}
    for k, v in mapping.items():
        try:
            col_idx = int(k)
        except (ValueError, TypeError):
            raise PollValidationError(f"Invalid column index in mapping: '{k}'.")
        if not (0 <= col_idx < num_cols):
            raise PollValidationError(f"Column index {col_idx} out of range (0..{num_cols - 1}).")
        col_to_target[col_idx] = v.strip()

    name_cols = [c for c, t in col_to_target.items() if t == "name"]
    if len(name_cols) == 0:
        raise PollValidationError("Exactly one column must be mapped to 'name'.")
    if len(name_cols) > 1:
        raise PollValidationError("Only one column may be mapped to 'name'.")
    name_col_idx = name_cols[0]

    # Check that no two columns map to the same target (other than 'skip')
    mapped_targets = [t for t in col_to_target.values() if t != "skip"]
    target_counts = Counter(mapped_targets)
    duplicates = [t for t, count in target_counts.items() if count > 1]
    if duplicates:
        raise PollValidationError("Two columns may not map to the same target.")

    # 3. Create new fields if specified (within this transaction)
    fields = list_group_fields(session, group.id)
    existing_field_keys = {f.key for f in fields}
    fields_created_names: list[str] = []

    if new_fields_data:
        for nf in new_fields_data:
            col_raw = nf.get("column")
            if col_raw is None:
                raise PollValidationError("Each new_fields entry must specify a 'column'.")

            # Resolve column index
            if isinstance(col_raw, int) or (isinstance(col_raw, str) and col_raw.isdigit()):
                nf_col = int(col_raw)
            else:
                # Match by header name
                try:
                    nf_col = headers.index(str(col_raw).strip())
                except ValueError:
                    raise PollValidationError(f"Column '{col_raw}' in new_fields not found in headers.")

            if not (0 <= nf_col < num_cols):
                raise PollValidationError(f"Column index {nf_col} in new_fields out of range.")

            # Rule: new_fields entry pointing at a column that is also mapped is an error
            current_mapping = col_to_target.get(nf_col, "skip")
            if current_mapping != "skip":
                raise PollValidationError(
                    f"A new_fields entry points to column {nf_col} which is already mapped to '{current_mapping}'."
                )

            nf_name = nf.get("name")
            if not nf_name or not isinstance(nf_name, str):
                raise PollValidationError("New field must have a non-empty 'name'.")

            created_f = create_group_field(
                session,
                group_id=group.id,
                name=nf_name,
                field_type=nf.get("field_type", "text"),
                is_required=bool(nf.get("is_required", False)),
                default_value=nf.get("default_value"),
                choices=nf.get("choices"),
                is_identifier=bool(nf.get("is_identifier", False)),
                commit=False,
            )
            fields.append(created_f)
            fields_created_names.append(created_f.name)
            existing_field_keys.add(created_f.key)
            col_to_target[nf_col] = created_f.key

    # Check for unknown field keys in mapping
    for t in mapped_targets:
        if t != "name" and t not in existing_field_keys:
            raise PollValidationError(f"Unknown field key: '{t}'.")

    # Map field keys to GroupField objects
    field_map = {f.key: f for f in fields}
    id_field = next((f for f in fields if f.is_identifier), None)

    # 4. Fetch existing members for duplicate checks
    existing_members = list(
        session.scalars(select(Member).where(Member.group_id == group.id)).all()
    )
    existing_names: set[str] = {m.display_name for m in existing_members}
    existing_identifiers: set[str] = {
        m.identifier_value for m in existing_members if m.identifier_value is not None
    }

    # 5. Row-by-row validation
    errors: list[dict[str, Any]] = []
    seen_identifiers_in_file: dict[str, int] = {}
    seen_names_in_file: dict[str, int] = {}

    rows_total = len(data_rows)
    rows_skipped = 0
    members_to_add_data: list[dict[str, Any]] = []

    for sheet_row_num, cells in data_rows:
        row_has_error = False

        # Display name
        raw_name = cells[name_col_idx]
        conv_name = convert_cell_value(raw_name)
        if is_blank(conv_name):
            errors.append({
                "row": sheet_row_num,
                "field": "Name",
                "message": "Member display_name cannot be blank.",
            })
            row_has_error = True
            clean_name = ""
        elif len(conv_name) > 255:
            errors.append({
                "row": sheet_row_num,
                "field": "Name",
                "message": "Member display_name cannot exceed 255 characters.",
            })
            row_has_error = True
            clean_name = conv_name[:255]
        else:
            clean_name = conv_name

        # Extract values for mapped columns
        raw_values: dict[str, Any] = {}
        for c_idx, target in col_to_target.items():
            if target in ("name", "skip"):
                continue
            f_obj = field_map.get(target)
            f_type = f_obj.field_type if f_obj else None
            cell_val = cells[c_idx] if c_idx < len(cells) else None
            converted = convert_cell_value(cell_val, target_field_type=f_type)
            raw_values[target] = converted

        # Validate member values using R1a rules
        cleaned_vals: dict[str, Any] = {}
        identifier_val: str | None = None
        try:
            cleaned_vals, identifier_val = validate_member_values_for_create(
                fields, raw_values, row=sheet_row_num
            )
        except PollValidationError as exc:
            row_has_error = True
            if exc.details:
                errors.extend(exc.details)
            else:
                errors.append({
                    "row": sheet_row_num,
                    "field": "values",
                    "message": exc.message,
                })

        # Duplication checks
        is_row_skipped = False
        if not row_has_error:
            if id_field is not None:
                raw_id = raw_values.get(id_field.key, "")
                if identifier_val is not None:
                    # In-file duplicate check
                    if identifier_val in seen_identifiers_in_file:
                        earlier_row = seen_identifiers_in_file[identifier_val]
                        errors.append({
                            "row": sheet_row_num,
                            "field": id_field.name,
                            "message": f"Duplicate identifier '{raw_id}' (first seen on row {earlier_row}).",
                        })
                        row_has_error = True
                    else:
                        seen_identifiers_in_file[identifier_val] = sheet_row_num

                    # Existing member duplicate check
                    if not row_has_error and identifier_val in existing_identifiers:
                        if on_duplicate == "reject":
                            errors.append({
                                "row": sheet_row_num,
                                "field": id_field.name,
                                "message": f"Member with identifier '{raw_id}' already exists in group.",
                            })
                            row_has_error = True
                        else:
                            is_row_skipped = True
            else:
                # Group has NO identifier field: display name must be unique
                if clean_name:
                    if clean_name in seen_names_in_file:
                        earlier_row = seen_names_in_file[clean_name]
                        errors.append({
                            "row": sheet_row_num,
                            "field": "Name",
                            "message": f"Duplicate name '{clean_name}' (first seen on row {earlier_row}).",
                        })
                        row_has_error = True
                    else:
                        seen_names_in_file[clean_name] = sheet_row_num

                    if not row_has_error and clean_name in existing_names:
                        if on_duplicate == "reject":
                            errors.append({
                                "row": sheet_row_num,
                                "field": "Name",
                                "message": f"Member with name '{clean_name}' already exists in group.",
                            })
                            row_has_error = True
                        else:
                            is_row_skipped = True

        if is_row_skipped:
            rows_skipped += 1
        elif not row_has_error:
            members_to_add_data.append({
                "display_name": clean_name,
                "values": cleaned_vals,
            })

    # 6. If any error occurred, abort and report
    if errors:
        session.rollback()
        err_count = len(errors)
        top_msg = (
            f"Import failed with {err_count} errors. Please fix them and try again."
            if err_count > 1
            else errors[0]["message"]
        )
        raise PollValidationError(top_msg, details=errors[:MAX_ERROR_DETAILS])

    rows_added = len(members_to_add_data)

    # 7. Dry run check
    if dry_run:
        session.rollback()
        return {
            "dry_run": True,
            "rows_total": rows_total,
            "rows_added": rows_added,
            "rows_skipped": rows_skipped,
            "fields_created": fields_created_names,
            "errors": [],
        }

    # 8. Save members
    try:
        if rows_added > 0:
            add_members_bulk(session, group.id, members_data=members_to_add_data)
        else:
            # Commit newly created fields even if 0 members were added
            session.commit()
    except Exception:
        session.rollback()
        raise

    return {
        "dry_run": False,
        "rows_total": rows_total,
        "rows_added": rows_added,
        "rows_skipped": rows_skipped,
        "fields_created": fields_created_names,
        "errors": [],
    }
