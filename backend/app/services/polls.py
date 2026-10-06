"""Polls service layer implementing all domain rules and operations."""

from collections.abc import Sequence
import csv
from dataclasses import dataclass
from datetime import datetime, timezone
import io
import re
from typing import Any
import uuid

import openpyxl
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.errors import (
    ClaimNotApprovedError,
    GroupNotFoundError,
    MemberGroupMismatchError,
    MemberInactiveError,
    MemberNotFoundError,
    OptionNotFoundError,
    OptionPollMismatchError,
    PollClosedError,
    PollNotFoundError,
    PollValidationError,
)
from app.models import (
    Group,
    GroupField,
    Member,
    Poll,
    PollAnswer,
    PollField,
    PollIncludedField,
    PollOption,
    Vote,
    VoteHistory,
)
from app.security import generate_join_code
from app.services.fields import (
    generate_field_key,
    is_blank,
    validate_answers_for_update,
    validate_field_definition,
    validate_member_values_for_create,
)

VALID_ROLES = frozenset({"target", "in_progress", "excused", "not_yet"})
VALID_COMPLETION_MODES = frozenset({"first", "last"})
VALID_POLL_STATUSES = frozenset({"open", "closed"})


def ensure_utc(dt: datetime | None) -> datetime | None:
    """Ensure a datetime is timezone-aware UTC."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def resolve_uuid(val: uuid.UUID | str) -> uuid.UUID:
    """Resolve a UUID from either UUID or string."""
    if isinstance(val, uuid.UUID):
        return val
    return uuid.UUID(str(val))


@dataclass(frozen=True)
class PollOptionInput:
    """Input payload for a poll option."""

    label: str
    role: str


@dataclass
class AtTargetMember:
    """Member status entry for at_target category."""

    member: Member
    late: bool
    completed_at: datetime


@dataclass
class PollStatusResult:
    """Partitioned active member status lists for a poll."""

    at_target: list[AtTargetMember]
    excused: list[Member]
    behind_target: list[Member]
    not_voted: list[Member]


@dataclass
class MemberOptionHistory:
    """Historical selection details for a member and option."""

    option_id: uuid.UUID
    option_label: str
    first_selected_at: datetime
    last_selected_at: datetime
    is_selected: bool


# ---------------------------------------------------------------------------
# Group and Member Management
# ---------------------------------------------------------------------------


def create_group(
    session: Session,
    name: str,
    now: datetime | None = None,
) -> Group:
    """Create a new group."""
    if not name or not name.strip():
        raise PollValidationError("Group name cannot be empty.")

    current_time = ensure_utc(now) or datetime.now(timezone.utc)
    group = Group(
        id=uuid.uuid4(),
        name=name.strip(),
        join_code=generate_join_code(),
        admin_token_hash=None,
        require_claim_approval=False,
        created_at=current_time,
    )
    session.add(group)
    session.commit()
    session.refresh(group)
    return group


def add_member(
    session: Session,
    group_id: uuid.UUID | str,
    display_name: str,
    is_active: bool = True,
    now: datetime | None = None,
    values: dict[str, Any] | None = None,
) -> Member:
    """Add a new member to a group."""
    gid = resolve_uuid(group_id)

    # Concurrency lock
    group = session.scalar(
        select(Group).where(Group.id == gid).with_for_update()
    )
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    if not display_name or not display_name.strip():
        raise PollValidationError("Member display_name cannot be empty.")
    clean_name = display_name.strip()

    fields = list(
        session.scalars(
            select(GroupField)
            .where(GroupField.group_id == gid)
            .order_by(GroupField.position)
        ).all()
    )
    id_field = next((f for f in fields if f.is_identifier), None)

    cleaned_vals, identifier_val = validate_member_values_for_create(fields, values)

    if id_field is not None:
        if identifier_val is not None:
            existing = session.scalar(
                select(Member).where(
                    Member.group_id == gid,
                    Member.identifier_value == identifier_val,
                )
            )
            if existing:
                raise PollValidationError(
                    f"Member with identifier '{identifier_val}' already exists."
                )
    else:
        existing = session.scalar(
            select(Member).where(
                Member.group_id == gid,
                Member.display_name == clean_name,
            )
        )
        if existing:
            raise PollValidationError(
                f"Member with name '{clean_name}' already exists in group."
            )

    current_time = ensure_utc(now) or datetime.now(timezone.utc)
    member = Member(
        id=uuid.uuid4(),
        group_id=gid,
        display_name=clean_name,
        is_active=is_active,
        field_values=cleaned_vals,
        identifier_value=identifier_val,
        created_at=current_time,
    )
    if id_field is not None:
        raw_id = cleaned_vals.get(id_field.key)
        member.identifier = str(raw_id) if raw_id is not None else None
    else:
        member.identifier = None

    session.add(member)
    session.commit()
    session.refresh(member)
    if id_field is not None:
        raw_id = (member.field_values or {}).get(id_field.key)
        member.identifier = str(raw_id) if raw_id is not None else None
    else:
        member.identifier = None
    return member


# ---------------------------------------------------------------------------
# Poll Creation and Management
# ---------------------------------------------------------------------------


def create_poll(
    session: Session,
    group_id: uuid.UUID | str,
    name: str,
    allow_multiple: bool,
    options: Sequence[PollOptionInput | dict[str, Any] | tuple[str, str]],
    description_raw: str | None = None,
    description_final: str | None = None,
    deadline: datetime | None = None,
    completion_time_mode: str = "last",
    now: datetime | None = None,
    included_field_ids: Sequence[uuid.UUID | str] | None = None,
    poll_fields: Sequence[dict[str, Any]] | None = None,
) -> Poll:
    """Create a new poll with validated options, included group fields, and poll-only fields."""
    gid = resolve_uuid(group_id)
    group = session.get(Group, gid)
    if not group:
        raise GroupNotFoundError(f"Group {gid} not found.")

    if not name or not name.strip():
        raise PollValidationError("Poll name cannot be empty.")

    if len(options) < 2:
        raise PollValidationError("A poll must have at least 2 options.")

    if completion_time_mode not in VALID_COMPLETION_MODES:
        raise PollValidationError(
            f"Invalid completion_time_mode: '{completion_time_mode}'. Expected one of {list(VALID_COMPLETION_MODES)}."
        )

    # 1. Fetch group fields and identifier
    all_group_fields = list(
        session.scalars(
            select(GroupField)
            .where(GroupField.group_id == gid)
            .order_by(GroupField.position)
        ).all()
    )
    group_field_map = {f.id: f for f in all_group_fields}
    identifier_field = next((f for f in all_group_fields if f.is_identifier), None)

    # 2. Included group fields validation
    seen_included_ids: set[uuid.UUID] = set()
    ordered_included_fields: list[GroupField] = []
    if included_field_ids:
        for raw_fid in included_field_ids:
            fid = resolve_uuid(raw_fid)
            if fid not in group_field_map:
                raise PollValidationError(f"Field {fid} does not belong to group {gid}.")
            if fid in seen_included_ids:
                continue
            seen_included_ids.add(fid)
            gf = group_field_map[fid]
            # Rule: The identifier field is not stored here; it is added automatically when results are built.
            if identifier_field and gf.id == identifier_field.id:
                continue
            ordered_included_fields.append(gf)

    # 3. Poll-only fields validation
    validated_poll_fields: list[dict[str, Any]] = []
    if poll_fields:
        if len(poll_fields) > 15:
            raise PollValidationError("A poll can have at most 15 poll-only fields.")

        seen_poll_names: set[str] = set()
        existing_poll_keys: set[str] = set()

        for pf_spec in poll_fields:
            raw_name = pf_spec.get("name", "")
            ftype = pf_spec.get("field_type", "")
            is_req = bool(pf_spec.get("is_required", False))
            dval = pf_spec.get("default_value")
            choices = pf_spec.get("choices")

            clean_name, clean_choices, clean_default = validate_field_definition(
                raw_name,
                ftype,
                choices=choices,
                default_value=dval,
                is_required=is_req,
                is_identifier=False,
            )

            name_cf = clean_name.casefold()
            if name_cf in seen_poll_names:
                raise PollValidationError(f"Duplicate poll field name: '{clean_name}'.")
            if identifier_field and name_cf == identifier_field.name.casefold():
                raise PollValidationError(
                    f"Poll field name '{clean_name}' conflicts with group identifier field."
                )
            if any(name_cf == gf.name.casefold() for gf in ordered_included_fields):
                raise PollValidationError(
                    f"Poll field name '{clean_name}' conflicts with included group field."
                )

            seen_poll_names.add(name_cf)
            key = generate_field_key(clean_name, existing_poll_keys)
            existing_poll_keys.add(key)

            validated_poll_fields.append({
                "key": key,
                "name": clean_name,
                "field_type": ftype,
                "is_required": is_req,
                "default_value": clean_default,
                "choices": clean_choices,
            })

    # Convert options to PollOptionInput objects
    normalized_options: list[PollOptionInput] = []
    seen_labels: set[str] = set()
    has_target = False

    for opt in options:
        if isinstance(opt, PollOptionInput):
            label, role = opt.label, opt.role
        elif isinstance(opt, dict):
            label, role = opt["label"], opt["role"]
        elif isinstance(opt, (tuple, list)) and len(opt) == 2:
            label, role = opt[0], opt[1]
        else:
            raise PollValidationError(f"Invalid option specification: {opt}")

        label = str(label)
        role = str(role)

        if not label:
            raise PollValidationError("Poll option label cannot be empty.")

        if label in seen_labels:
            raise PollValidationError(
                f"Duplicate option label detected: '{label}'. Labels must be unique within a poll."
            )
        seen_labels.add(label)

        if role not in VALID_ROLES:
            raise PollValidationError(
                f"Invalid role: '{role}'. Expected one of {list(VALID_ROLES)}."
            )

        if role == "target":
            has_target = True

        normalized_options.append(PollOptionInput(label=label, role=role))

    if not has_target:
        raise PollValidationError(
            "A poll must have at least one option with role 'target'."
        )

    current_time = ensure_utc(now) or datetime.now(timezone.utc)
    deadline_utc = ensure_utc(deadline)

    poll = Poll(
        id=uuid.uuid4(),
        group_id=gid,
        name=name.strip(),
        description_raw=description_raw,
        description_final=description_final,
        status="open",
        allow_multiple=allow_multiple,
        deadline=deadline_utc,
        completion_time_mode=completion_time_mode,
        closed_at=None,
        created_at=current_time,
    )
    session.add(poll)
    session.flush()

    # Assign position in the order given (1-based index)
    for position, opt_input in enumerate(normalized_options, start=1):
        option = PollOption(
            id=uuid.uuid4(),
            poll_id=poll.id,
            label=opt_input.label,
            position=position,
            role=opt_input.role,
            created_at=current_time,
        )
        session.add(option)

    # Add included fields (position 1-based)
    for position, gf in enumerate(ordered_included_fields, start=1):
        pif = PollIncludedField(
            poll_id=poll.id,
            field_id=gf.id,
            position=position,
        )
        session.add(pif)

    # Add poll-only fields (position 1-based)
    for position, pf_data in enumerate(validated_poll_fields, start=1):
        pf = PollField(
            id=uuid.uuid4(),
            poll_id=poll.id,
            key=pf_data["key"],
            name=pf_data["name"],
            field_type=pf_data["field_type"],
            is_required=pf_data["is_required"],
            default_value=pf_data["default_value"],
            choices=pf_data["choices"],
            position=position,
            created_at=current_time,
        )
        session.add(pf)

    session.commit()
    session.refresh(poll)
    return poll


def close_poll(
    session: Session,
    poll_id: uuid.UUID | str,
    now: datetime | None = None,
) -> Poll:
    """Close an open poll. Closing an already closed poll is a no-op."""
    pid = resolve_uuid(poll_id)
    poll = session.get(Poll, pid)
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    if poll.status == "closed":
        return poll

    close_time = ensure_utc(now) or datetime.now(timezone.utc)
    poll.status = "closed"
    poll.closed_at = close_time
    session.commit()
    session.refresh(poll)
    return poll


# ---------------------------------------------------------------------------
# Voting Operations
# ---------------------------------------------------------------------------


def cast_vote(
    session: Session,
    poll_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
    option_id: uuid.UUID | str,
    now: datetime | None = None,
) -> None:
    """Cast or switch a member's vote for an option in a poll."""
    pid = resolve_uuid(poll_id)
    mid = resolve_uuid(member_id)
    oid = resolve_uuid(option_id)

    poll = session.scalar(
        select(Poll).where(Poll.id == pid).with_for_update().execution_options(populate_existing=True)
    )
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    if poll.status == "closed":
        raise PollClosedError(f"Poll {pid} is closed. Cannot cast votes.")

    member = session.get(Member, mid)
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != poll.group_id:
        raise MemberGroupMismatchError(
            f"Member {mid} does not belong to poll's group {poll.group_id}."
        )

    if not member.is_active:
        raise MemberInactiveError(f"Member {mid} is inactive.")

    option = session.get(PollOption, oid)
    if not option:
        raise OptionNotFoundError(f"Option {oid} not found.")

    if option.poll_id != poll.id:
        raise OptionPollMismatchError(
            f"Option {oid} does not belong to poll {poll.id}."
        )

    current_time = ensure_utc(now) or datetime.now(timezone.utc)

    # Check if option is already currently selected by this member
    existing_selection = session.scalar(
        select(Vote).where(
            Vote.poll_id == pid,
            Vote.member_id == mid,
            Vote.option_id == oid,
        )
    )
    if existing_selection is not None:
        # Re-selecting an already selected option: do nothing (no timestamp change)
        return

    # If single-choice poll, delete member's other current selections in this poll
    if not poll.allow_multiple:
        session.execute(
            delete(Vote).where(
                Vote.poll_id == pid,
                Vote.member_id == mid,
            )
        )

    # Add new selection row
    new_vote = Vote(
        id=uuid.uuid4(),
        poll_id=pid,
        member_id=mid,
        option_id=oid,
        created_at=current_time,
    )
    session.add(new_vote)

    # Update vote_history for this (poll_id, member_id, option_id)
    history_entry = session.scalar(
        select(VoteHistory).where(
            VoteHistory.poll_id == pid,
            VoteHistory.member_id == mid,
            VoteHistory.option_id == oid,
        )
    )

    if history_entry is None:
        # Create history entry
        new_history = VoteHistory(
            id=uuid.uuid4(),
            poll_id=pid,
            member_id=mid,
            option_id=oid,
            first_selected_at=current_time,
            last_selected_at=current_time,
            created_at=current_time,
        )
        session.add(new_history)
    else:
        # Update only last_selected_at
        history_entry.last_selected_at = current_time

    session.commit()


def remove_vote(
    session: Session,
    poll_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
    option_id: uuid.UUID | str | None = None,
    now: datetime | None = None,
) -> None:
    """Remove a specific selection or all selections for a member."""
    pid = resolve_uuid(poll_id)
    mid = resolve_uuid(member_id)

    poll = session.scalar(
        select(Poll).where(Poll.id == pid).with_for_update().execution_options(populate_existing=True)
    )
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    if poll.status == "closed":
        raise PollClosedError(f"Poll {pid} is closed. Cannot remove votes.")

    member = session.get(Member, mid)
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != poll.group_id:
        raise MemberGroupMismatchError(
            f"Member {mid} does not belong to poll's group {poll.group_id}."
        )

    if option_id is not None:
        oid = resolve_uuid(option_id)
        option = session.get(PollOption, oid)
        if not option:
            raise OptionNotFoundError(f"Option {oid} not found.")

        if option.poll_id != poll.id:
            raise OptionPollMismatchError(
                f"Option {oid} does not belong to poll {poll.id}."
            )

        # Remove only that selection
        session.execute(
            delete(Vote).where(
                Vote.poll_id == pid,
                Vote.member_id == mid,
                Vote.option_id == oid,
            )
        )
    else:
        # Remove all current selections in this poll
        session.execute(
            delete(Vote).where(
                Vote.poll_id == pid,
                Vote.member_id == mid,
            )
        )

    # vote_history is never touched
    session.commit()


# ---------------------------------------------------------------------------
# Status and Reporting Queries
# ---------------------------------------------------------------------------


def get_poll_status(
    session: Session,
    poll_id: uuid.UUID | str,
    now: datetime | None = None,
) -> PollStatusResult:
    """Calculate and return the 4 status categories for all active members in the poll's group."""
    pid = resolve_uuid(poll_id)
    poll = session.get(Poll, pid)
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    # Get all active members of the poll's group
    active_members = (
        session.scalars(
            select(Member)
            .where(
                Member.group_id == poll.group_id,
                Member.is_active.is_(True),
            )
            .order_by(Member.display_name)
        )
        .all()
    )

    # Fetch all current votes for this poll with option and history details
    votes = session.scalars(
        select(Vote).where(Vote.poll_id == pid)
    ).all()

    # Map member_id -> list of Vote objects
    member_votes: dict[uuid.UUID, list[Vote]] = {}
    for v in votes:
        member_votes.setdefault(v.member_id, []).append(v)

    # Fetch all options for this poll
    options_by_id = {
        opt.id: opt
        for opt in session.scalars(
            select(PollOption).where(PollOption.poll_id == pid)
        ).all()
    }

    # Fetch all history rows for this poll
    history_rows = session.scalars(
        select(VoteHistory).where(VoteHistory.poll_id == pid)
    ).all()
    history_by_key = {
        (h.member_id, h.option_id): h for h in history_rows
    }

    at_target: list[AtTargetMember] = []
    excused: list[Member] = []
    behind_target: list[Member] = []
    not_voted: list[Member] = []

    for member in active_members:
        cur_votes = member_votes.get(member.id, [])
        if not cur_votes:
            not_voted.append(member)
            continue

        selected_options = [
            options_by_id[v.option_id]
            for v in cur_votes
            if v.option_id in options_by_id
        ]
        target_options = [opt for opt in selected_options if opt.role == "target"]
        excused_options = [opt for opt in selected_options if opt.role == "excused"]
        behind_options = [
            opt for opt in selected_options if opt.role in ("in_progress", "not_yet")
        ]

        if target_options:
            # Rule 6 & 7: at_target
            # Calculate completed_at and late
            candidate_times: list[datetime] = []
            for opt in target_options:
                vh = history_by_key.get((member.id, opt.id))
                if vh is not None:
                    if poll.completion_time_mode == "first":
                        candidate_times.append(vh.first_selected_at)
                    else:
                        candidate_times.append(vh.last_selected_at)

            if not candidate_times:
                # Fallback to vote created_at if history record is missing
                candidate_times = [
                    v.created_at for v in cur_votes if v.option_id in {o.id for o in target_options}
                ]

            completed_at = min(candidate_times)
            late = (poll.deadline is not None) and (completed_at > poll.deadline)
            at_target.append(
                AtTargetMember(
                    member=member,
                    late=late,
                    completed_at=completed_at,
                )
            )
        elif excused_options:
            # Rule 6: excused
            excused.append(member)
        elif behind_options:
            # Rule 6: behind_target
            behind_target.append(member)
        else:
            not_voted.append(member)

    return PollStatusResult(
        at_target=at_target,
        excused=excused,
        behind_target=behind_target,
        not_voted=not_voted,
    )


def all_reached(
    session: Session,
    poll_id: uuid.UUID | str,
    now: datetime | None = None,
) -> bool:
    """Return True if every active member is at target or excused, and at least one is at target."""
    status = get_poll_status(session, poll_id, now=now)
    total_active_members = (
        len(status.at_target)
        + len(status.excused)
        + len(status.behind_target)
        + len(status.not_voted)
    )

    if total_active_members == 0:
        return False

    if len(status.at_target) == 0:
        return False

    if len(status.behind_target) > 0 or len(status.not_voted) > 0:
        return False

    return True


def get_member_history(
    session: Session,
    poll_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
) -> list[MemberOptionHistory]:
    """Return chronological history of all options ever selected by a member in a poll."""
    pid = resolve_uuid(poll_id)
    mid = resolve_uuid(member_id)

    poll = session.get(Poll, pid)
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    member = session.get(Member, mid)
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    # Current selections for this member in this poll
    current_selected_option_ids = set(
        session.scalars(
            select(Vote.option_id).where(
                Vote.poll_id == pid,
                Vote.member_id == mid,
            )
        ).all()
    )

    # History rows joined with options
    history_entries = (
        session.scalars(
            select(VoteHistory)
            .where(
                VoteHistory.poll_id == pid,
                VoteHistory.member_id == mid,
            )
            .order_by(VoteHistory.first_selected_at)
        )
        .all()
    )

    options_by_id = {
        opt.id: opt
        for opt in session.scalars(
            select(PollOption).where(PollOption.poll_id == pid)
        ).all()
    }

    result: list[MemberOptionHistory] = []
    for h in history_entries:
        opt = options_by_id.get(h.option_id)
        label = opt.label if opt else str(h.option_id)
        is_selected = h.option_id in current_selected_option_ids
        result.append(
            MemberOptionHistory(
                option_id=h.option_id,
                option_label=label,
                first_selected_at=h.first_selected_at,
                last_selected_at=h.last_selected_at,
                is_selected=is_selected,
            )
        )

    return result


# ---------------------------------------------------------------------------
# Poll Answers Management
# ---------------------------------------------------------------------------


def save_poll_answers(
    session: Session,
    poll_id: uuid.UUID | str,
    member_id: uuid.UUID | str,
    values: dict[str, Any],
    now: datetime | None = None,
) -> PollAnswer:
    """Save or update member answers for a poll with row locking on the poll."""
    pid = resolve_uuid(poll_id)
    mid = resolve_uuid(member_id)
    current_time = ensure_utc(now) or datetime.now(timezone.utc)

    # Concurrency lock on poll row with populate_existing
    poll = session.scalar(
        select(Poll)
        .where(Poll.id == pid)
        .with_for_update()
        .execution_options(populate_existing=True)
    )
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    if poll.status == "closed":
        raise PollClosedError("Cannot answer a closed poll.")

    member = session.get(Member, mid)
    if not member:
        raise MemberNotFoundError(f"Member {mid} not found.")

    if member.group_id != poll.group_id:
        raise MemberGroupMismatchError("This poll belongs to a different group.")

    if member.claim_status != "approved":
        raise ClaimNotApprovedError("Member claim has not been approved.")

    if not member.is_active:
        raise MemberInactiveError("Member is inactive.")

    poll_fields = list(
        session.scalars(
            select(PollField)
            .where(PollField.poll_id == poll.id)
            .order_by(PollField.position)
        ).all()
    )

    answer = session.scalar(
        select(PollAnswer)
        .where(PollAnswer.poll_id == poll.id, PollAnswer.member_id == member.id)
        .with_for_update()
    )

    current_answers = dict(answer.values) if answer and answer.values else {}
    updated_answers = validate_answers_for_update(poll_fields, current_answers, values)

    if answer is None:
        answer = PollAnswer(
            id=uuid.uuid4(),
            poll_id=poll.id,
            member_id=member.id,
            values=dict(updated_answers),  # Assign a fresh dict
            updated_at=current_time,
        )
        session.add(answer)
    else:
        answer.values = dict(updated_answers)  # Assign a fresh dict
        answer.updated_at = current_time

    session.commit()
    session.refresh(answer)
    return answer


# ---------------------------------------------------------------------------
# Poll Results Table & Export
# ---------------------------------------------------------------------------


def get_poll_results(
    session: Session,
    poll_id: uuid.UUID | str,
) -> dict[str, Any]:
    """Build poll results table data for admin view."""
    pid = resolve_uuid(poll_id)
    poll = session.get(Poll, pid)
    if not poll:
        raise PollNotFoundError(f"Poll {pid} not found.")

    group = session.get(Group, poll.group_id)
    if not group:
        raise GroupNotFoundError(f"Group {poll.group_id} not found.")

    # 1. Identifier field (if any)
    id_field = session.scalar(
        select(GroupField).where(
            GroupField.group_id == group.id,
            GroupField.is_identifier.is_(True),
        )
    )

    # 2. Included group fields in position order (skipping identifier if present)
    included_assocs = list(
        session.scalars(
            select(PollIncludedField)
            .where(PollIncludedField.poll_id == poll.id)
            .order_by(PollIncludedField.position)
        ).all()
    )
    included_group_fields: list[GroupField] = []
    for assoc in included_assocs:
        f = assoc.field
        if f is not None:
            if id_field and f.id == id_field.id:
                continue
            included_group_fields.append(f)

    # 3. Poll-only fields in position order
    poll_fields = list(
        session.scalars(
            select(PollField)
            .where(PollField.poll_id == poll.id)
            .order_by(PollField.position)
        ).all()
    )

    # Build columns list
    columns: list[dict[str, Any]] = []
    if id_field:
        columns.append({
            "source": "group",
            "key": id_field.key,
            "name": id_field.name,
            "field_type": id_field.field_type,
            "is_identifier": True,
        })
    for gf in included_group_fields:
        columns.append({
            "source": "group",
            "key": gf.key,
            "name": gf.name,
            "field_type": gf.field_type,
            "is_identifier": False,
        })
    for pf in poll_fields:
        columns.append({
            "source": "poll",
            "key": pf.key,
            "name": pf.name,
            "field_type": pf.field_type,
            "is_identifier": False,
        })

    # Active members only
    active_members = list(
        session.scalars(
            select(Member)
            .where(Member.group_id == group.id, Member.is_active.is_(True))
            .order_by(Member.display_name, Member.created_at, Member.id)
        ).all()
    )

    # Status mapping from get_poll_status
    status_result = get_poll_status(session, poll.id)
    status_map: dict[uuid.UUID, tuple[str, bool | None, datetime | None]] = {}
    for item in status_result.at_target:
        status_map[item.member.id] = ("at_target", item.late, item.completed_at)
    for m in status_result.excused:
        status_map[m.id] = ("excused", None, None)
    for m in status_result.behind_target:
        status_map[m.id] = ("behind_target", None, None)
    for m in status_result.not_voted:
        status_map[m.id] = ("not_voted", None, None)

    # Current votes mapping: member_id -> list of selected option labels (position order)
    votes = list(
        session.scalars(
            select(Vote).where(Vote.poll_id == poll.id)
        ).all()
    )
    votes_sorted = sorted(votes, key=lambda v: v.option.position if v.option else 0)
    member_votes: dict[uuid.UUID, list[str]] = {m.id: [] for m in active_members}
    for v in votes_sorted:
        if v.member_id in member_votes and v.option:
            member_votes[v.member_id].append(v.option.label)

    # Member answers mapping
    answers = list(
        session.scalars(
            select(PollAnswer).where(PollAnswer.poll_id == poll.id)
        ).all()
    )
    answers_map: dict[uuid.UUID, PollAnswer] = {a.member_id: a for a in answers}

    # Required poll fields for answers_complete calculation
    required_poll_fields = [pf for pf in poll_fields if pf.is_required]

    rows: list[dict[str, Any]] = []
    group_keys = ([id_field.key] if id_field else []) + [gf.key for gf in included_group_fields]

    for m in active_members:
        st_info = status_map.get(m.id, ("not_voted", None, None))
        status_name, late, completed_at = st_info

        ans_obj = answers_map.get(m.id)
        member_ans = dict(ans_obj.values) if ans_obj and ans_obj.values else {}
        ans_updated_at = ans_obj.updated_at if ans_obj else None

        # answers_complete calculation:
        # A member counts as answers_complete only when every required poll-only field has a non-blank saved value.
        # Members who never saved anything are incomplete if any field is required, and complete if no field is required.
        if not required_poll_fields:
            answers_complete = True
        else:
            if ans_obj is None or not ans_obj.values:
                answers_complete = False
            else:
                answers_complete = all(not is_blank(member_ans.get(pf.key)) for pf in required_poll_fields)

        # group_values dict (live read from m.field_values)
        m_field_values = m.field_values or {}
        g_vals: dict[str, Any] = {}
        for k in group_keys:
            if k in m_field_values:
                g_vals[k] = m_field_values[k]

        identifier_val: str | None = None
        if id_field and id_field.key in m_field_values:
            raw_v = m_field_values[id_field.key]
            identifier_val = str(raw_v) if not is_blank(raw_v) else None

        rows.append({
            "member_id": m.id,
            "display_name": m.display_name,
            "identifier": identifier_val,
            "status": status_name,
            "selected_options": member_votes.get(m.id, []),
            "late": late,
            "completed_at": completed_at,
            "group_values": g_vals,
            "answers": member_ans,
            "answers_updated_at": ans_updated_at,
            "answers_complete": answers_complete,
        })

    poll_info = {
        "id": poll.id,
        "name": poll.name,
        "status": poll.status,
        "deadline": poll.deadline,
        "allow_multiple": poll.allow_multiple,
    }

    return {
        "poll": poll_info,
        "columns": columns,
        "rows": rows,
    }


def make_safe_filename(poll_name: str, ext: str) -> str:
    """Generate safe filename for download from poll name."""
    safe = re.sub(r'[^a-zA-Z0-9_\-]+', '_', poll_name.strip())
    safe = safe.strip('_')
    if not safe:
        safe = "poll_results"
    return f"{safe[:100]}.{ext}"


def sanitize_formula_injection(val: Any) -> Any:
    """Prepend single quote to text cells starting with formula trigger chars (=, +, -, @)."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return str(val)
    if isinstance(val, (int, float)):
        return val
    s = str(val)
    stripped = s.strip()
    if stripped and stripped[0] in ("=", "+", "-", "@"):
        return f"'{s}"
    return s


def export_poll_results_csv(results: dict[str, Any]) -> tuple[bytes, str]:
    """Generate CSV bytes (UTF-8 with BOM) and safe filename for poll results."""
    poll_name = results["poll"]["name"]
    filename = make_safe_filename(poll_name, "csv")

    columns = results["columns"]
    headers = ["Name"] + [col["name"] for col in columns] + [
        "Status",
        "Selected options",
        "Late",
        "Completed at",
        "Answers complete",
    ]

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([sanitize_formula_injection(h) for h in headers])

    for row in results["rows"]:
        row_cells: list[Any] = [row["display_name"]]
        for col in columns:
            if col["source"] == "group":
                val = row["group_values"].get(col["key"], "")
            else:
                val = row["answers"].get(col["key"], "")
            row_cells.append(val)

        row_cells.append(row["status"])
        row_cells.append("; ".join(row["selected_options"]))
        if row["late"] is True:
            row_cells.append("Yes")
        elif row["late"] is False:
            row_cells.append("No")
        else:
            row_cells.append("")

        comp = row["completed_at"]
        if comp is not None:
            comp_dt = ensure_utc(comp)
            row_cells.append(comp_dt.isoformat())
        else:
            row_cells.append("")

        row_cells.append("Yes" if row["answers_complete"] else "No")
        writer.writerow([sanitize_formula_injection(cell) for cell in row_cells])

    csv_bytes = output.getvalue().encode("utf-8-sig")
    return csv_bytes, filename


def export_poll_results_xlsx(results: dict[str, Any]) -> tuple[bytes, str]:
    """Generate Excel (.xlsx) bytes and safe filename for poll results."""
    poll_name = results["poll"]["name"]
    filename = make_safe_filename(poll_name, "xlsx")

    columns = results["columns"]
    headers = ["Name"] + [col["name"] for col in columns] + [
        "Status",
        "Selected options",
        "Late",
        "Completed at",
        "Answers complete",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Poll Results"

    for c_idx, h in enumerate(headers, start=1):
        sanitized = sanitize_formula_injection(h)
        cell = ws.cell(row=1, column=c_idx, value=str(sanitized))
        cell.data_type = "s"

    for r_idx, row in enumerate(results["rows"], start=2):
        row_cells: list[Any] = [row["display_name"]]
        for col in columns:
            if col["source"] == "group":
                val = row["group_values"].get(col["key"], "")
            else:
                val = row["answers"].get(col["key"], "")
            row_cells.append(val)

        row_cells.append(row["status"])
        row_cells.append("; ".join(row["selected_options"]))
        if row["late"] is True:
            row_cells.append("Yes")
        elif row["late"] is False:
            row_cells.append("No")
        else:
            row_cells.append("")

        comp = row["completed_at"]
        if comp is not None:
            comp_dt = ensure_utc(comp)
            row_cells.append(comp_dt.isoformat())
        else:
            row_cells.append("")

        row_cells.append("Yes" if row["answers_complete"] else "No")

        for c_idx, val in enumerate(row_cells, start=1):
            sanitized = sanitize_formula_injection(val)
            cell = ws.cell(row=r_idx, column=c_idx)
            if isinstance(sanitized, (int, float)) and not isinstance(sanitized, bool):
                cell.value = sanitized
                cell.data_type = "n"
            else:
                cell.value = str(sanitized) if sanitized is not None else ""
                cell.data_type = "s"

    out_stream = io.BytesIO()
    wb.save(out_stream)
    return out_stream.getvalue(), filename
